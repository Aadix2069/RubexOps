import json
import os
import sys
from pathlib import Path

from openpyxl import load_workbook


# ============================================================
# CONFIG
# ============================================================

CONFIG_PATH = (
    Path.home()
    / "Documents"
    / "RubexOps"
    / "database_config.json"
)

SCON_SHEET_NAME = "scon"
START_ROW = 5
DEFAULT_ITEM_CODE = "NRFC"


# ============================================================
# ERROR HANDLER
# ============================================================

def fail(message):
    print(f"ERROR: {message}", file=sys.stderr)
    sys.exit(1)


# ============================================================
# HELPERS
# ============================================================

def get_database_path():
    if not CONFIG_PATH.exists():
        fail(f"Database config file not found: {CONFIG_PATH}")

    try:
        with open(CONFIG_PATH, "r", encoding="utf-8") as file:
            config = json.load(file)
    except json.JSONDecodeError:
        fail("Invalid JSON inside database_config.json.")

    database_path = (
        config.get("database_path")
        or config.get("excel_path")
        or config.get("path")
        or ""
    ).strip()

    if not database_path:
        fail("Database path was not found in database_config.json.")

    if not os.path.exists(database_path):
        fail(f"Excel database file not found: {database_path}")

    return database_path


def safe_string(value):
    if value is None:
        return ""

    return str(value).strip()


# ============================================================
# READ CUSTOMERS
# ============================================================

def read_customers():
    database_path = get_database_path()

    workbook = load_workbook(database_path, data_only=True, read_only=True)

    try:
        if SCON_SHEET_NAME not in workbook.sheetnames:
            fail(f"Sheet not found: {SCON_SHEET_NAME}")

        sheet = workbook[SCON_SHEET_NAME]
        customers = []
        seen = set()

        for row in range(START_ROW, sheet.max_row + 1):
            customer_name = safe_string(sheet[f"B{row}"].value)
            customer_id = safe_string(sheet[f"C{row}"].value)
            item_code = safe_string(sheet[f"E{row}"].value) or DEFAULT_ITEM_CODE
            renewal_reference = safe_string(sheet[f"Y{row}"].value)

            if not customer_name or not customer_id:
                continue

            if renewal_reference:
                continue

            key = f"{customer_id.casefold()}|{item_code.casefold()}"

            if key in seen:
                continue

            customers.append(
                {
                    "CustomerName": customer_name,
                    "CustomerID": customer_id,
                    "ItemCode": item_code,
                    "RenewalReference": renewal_reference,
                }
            )

            seen.add(key)

        return customers

    finally:
        workbook.close()


try:
    print(json.dumps(read_customers(), ensure_ascii=False))
except Exception as error:
    fail(str(error))
