import json
import os
import sys
from copy import copy
from datetime import datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.formula.translate import Translator


# ============================================================
# CONFIG
# ============================================================

CONFIG_PATH = (
    Path.home()
    / "Documents"
    / "RubexOps"
    / "database_config.json"
)

PURCHASE_SHEET_NAME = "purchase"
PRODUCTION_SHEET_NAME = "production"
START_ROW = 5


# ============================================================
# PRODUCTION COLUMN MAP
# ============================================================

COL_SERIAL_NO = "A"
COL_BATCH_ID = "B"
COL_PRODUCTION_DATE = "C"
COL_SUPPLIER_NAME = "D"
COL_SUPPLIER_ID = "E"
COL_INVOICE_NUMBER = "F"
COL_RAW_MATERIAL = "G"
COL_RAW_MATERIAL_CODE = "H"
COL_FINISHED_PRODUCT = "I"
COL_FINISHED_PRODUCT_CODE = "J"
COL_QUANTITY_AVAILABLE = "K"
COL_INPUT_WEIGHT = "L"
COL_OUTPUT_WEIGHT = "M"
COL_REMAINING_QUANTITY = "N"
COL_PRODUCTION_LOSS = "O"
COL_INITIAL_DRC = "P"
COL_ACTUAL_DRC = "Q"
COL_DRC_VARIANCE = "R"


# ============================================================
# PURCHASE COLUMN MAP
# ============================================================

PUR_COL_SUPPLIER_NAME = "B"
PUR_COL_SUPPLIER_ID = "C"
PUR_COL_RAW_MATERIAL = "D"
PUR_COL_RAW_MATERIAL_CODE = "E"
PUR_COL_INVOICE_NUMBER = "F"
PUR_COL_INVOICE_WEIGHT = "I"
PUR_COL_NET_WEIGHT = "N"
PUR_COL_DRC_WEIGHT = "P"


# ============================================================
# ERROR HANDLER
# ============================================================

def fail(message):
    print(f"ERROR: {message}", file=sys.stderr)
    sys.exit(1)


# ============================================================
# SAFE HELPERS
# ============================================================

def safe_string(value):
    if value is None:
        return ""

    return str(value).strip()


def parse_decimal(value, field_name):
    text = safe_string(value).replace(",", "").replace("%", "")

    if not text:
        fail(f"{field_name} is required.")

    try:
        return float(Decimal(text))
    except (InvalidOperation, ValueError):
        fail(f"{field_name} must be numeric.")


def number_or_zero(value):
    if value is None:
        return 0.0

    try:
        if isinstance(value, str):
            text = value.strip().replace(",", "").replace("%", "")

            if not text:
                return 0.0

            return float(Decimal(text))

        return float(value)
    except (InvalidOperation, ValueError, TypeError):
        return 0.0


def parse_date(value, field_name):
    text = safe_string(value)

    if not text:
        fail(f"{field_name} is required.")

    try:
        return datetime.strptime(text, "%d-%m-%Y")
    except ValueError:
        fail(f"{field_name} must be in dd-MM-yyyy format.")


def get_database_path():
    if not CONFIG_PATH.exists():
        fail(f"Database config file not found: {CONFIG_PATH}")

    try:
        with open(CONFIG_PATH, "r", encoding="utf-8") as file:
            config = json.load(file)
    except json.JSONDecodeError:
        fail("Invalid JSON inside database_config.json.")

    database_path = (
        config.get("database_path")
        or config.get("excel_path")
        or config.get("path")
        or ""
    ).strip()

    if not database_path:
        fail("Database path was not found in database_config.json.")

    if not os.path.exists(database_path):
        fail(f"Excel database file not found: {database_path}")

    return database_path


# ============================================================
# INVENTORY HELPERS
# ============================================================

def stock_key(supplier_id, invoice_number, raw_material_code):
    return (
        safe_string(supplier_id).casefold(),
        safe_string(invoice_number).casefold(),
        safe_string(raw_material_code).casefold(),
    )


def get_purchase_quantity(sheet, row):
    net_weight = number_or_zero(sheet[f"{PUR_COL_NET_WEIGHT}{row}"].value)

    if net_weight > 0:
        return net_weight

    drc_weight = number_or_zero(sheet[f"{PUR_COL_DRC_WEIGHT}{row}"].value)

    if drc_weight > 0:
        return drc_weight

    return number_or_zero(sheet[f"{PUR_COL_INVOICE_WEIGHT}{row}"].value)


def calculate_available_quantity(workbook, supplier_id, invoice_number, raw_material_code, exclude_row=None):
    if PURCHASE_SHEET_NAME not in workbook.sheetnames:
        fail(f"Sheet not found: {PURCHASE_SHEET_NAME}")

    key = stock_key(supplier_id, invoice_number, raw_material_code)
    purchase_sheet = workbook[PURCHASE_SHEET_NAME]
    purchased = 0.0

    for row in range(START_ROW, purchase_sheet.max_row + 1):
        row_key = stock_key(
            purchase_sheet[f"{PUR_COL_SUPPLIER_ID}{row}"].value,
            purchase_sheet[f"{PUR_COL_INVOICE_NUMBER}{row}"].value,
            purchase_sheet[f"{PUR_COL_RAW_MATERIAL_CODE}{row}"].value,
        )

        if row_key == key:
            purchased += get_purchase_quantity(purchase_sheet, row)

    consumed = 0.0

    if PRODUCTION_SHEET_NAME in workbook.sheetnames:
        production_sheet = workbook[PRODUCTION_SHEET_NAME]

        for row in range(START_ROW, production_sheet.max_row + 1):
            if exclude_row is not None and row == exclude_row:
                continue

            row_key = stock_key(
                production_sheet[f"{COL_SUPPLIER_ID}{row}"].value,
                production_sheet[f"{COL_INVOICE_NUMBER}{row}"].value,
                production_sheet[f"{COL_RAW_MATERIAL_CODE}{row}"].value,
            )

            if row_key == key:
                consumed += number_or_zero(production_sheet[f"{COL_INPUT_WEIGHT}{row}"].value)

    return round(max(purchased - consumed, 0.0), 4)


# ============================================================
# ROW HELPERS
# ============================================================

def generate_next_batch_id(sheet):
    highest = 0

    for row in range(START_ROW, sheet.max_row + 1):
        batch_id = safe_string(sheet[f"{COL_BATCH_ID}{row}"].value)

        if not batch_id.upper().startswith("BATCH-"):
            continue

        try:
            number = int(batch_id.split("-", 1)[1])
        except (IndexError, ValueError):
            continue

        highest = max(highest, number)

    return f"BATCH-{highest + 1:04d}"


def get_next_production_row(sheet):
    row = START_ROW

    while True:
        batch_id = sheet[f"{COL_BATCH_ID}{row}"].value
        invoice_number = sheet[f"{COL_INVOICE_NUMBER}{row}"].value

        if safe_string(batch_id) == "" and safe_string(invoice_number) == "":
            return row

        row += 1


def copy_row_style_and_formulas(sheet, source_row, target_row):
    for column in range(1, sheet.max_column + 1):
        source_cell = sheet.cell(row=source_row, column=column)
        target_cell = sheet.cell(row=target_row, column=column)

        if source_cell.has_style:
            target_cell.font = copy(source_cell.font)
            target_cell.fill = copy(source_cell.fill)
            target_cell.border = copy(source_cell.border)
            target_cell.alignment = copy(source_cell.alignment)
            target_cell.number_format = source_cell.number_format
            target_cell.protection = copy(source_cell.protection)

        if isinstance(source_cell.value, str) and source_cell.value.startswith("="):
            target_cell.value = Translator(
                source_cell.value,
                origin=source_cell.coordinate
            ).translate_formula(target_cell.coordinate)

    if source_row in sheet.row_dimensions:
        sheet.row_dimensions[target_row].height = sheet.row_dimensions[source_row].height


def write_calculation_formulas(sheet, row):
    sheet[f"{COL_REMAINING_QUANTITY}{row}"] = (
        f"={COL_QUANTITY_AVAILABLE}{row}-{COL_INPUT_WEIGHT}{row}"
    )
    sheet[f"{COL_PRODUCTION_LOSS}{row}"] = (
        f"={COL_INPUT_WEIGHT}{row}-{COL_OUTPUT_WEIGHT}{row}"
    )
    sheet[f"{COL_DRC_VARIANCE}{row}"] = (
        f"={COL_ACTUAL_DRC}{row}-{COL_INITIAL_DRC}{row}"
    )


# ============================================================
# CREATE ENTRY
# ============================================================

def create_production_entry(args):
    if len(args) != 12:
        fail(
            "Expected 12 arguments: production date, supplier name, supplier id, "
            "invoice number, raw material, raw material code, finished product, "
            "finished product code, input weight, output weight, initial DRC, actual DRC."
        )

    (
        production_date_text,
        supplier_name,
        supplier_id,
        invoice_number,
        raw_material,
        raw_material_code,
        finished_product,
        finished_product_code,
        input_weight_text,
        output_weight_text,
        initial_drc_text,
        actual_drc_text,
    ) = args

    production_date = parse_date(production_date_text, "Production Date")
    supplier_name = safe_string(supplier_name)
    supplier_id = safe_string(supplier_id)
    invoice_number = safe_string(invoice_number)
    raw_material = safe_string(raw_material)
    raw_material_code = safe_string(raw_material_code)
    finished_product = safe_string(finished_product)
    finished_product_code = safe_string(finished_product_code)

    if not supplier_name:
        fail("Supplier Name is required.")

    if not supplier_id:
        fail("Supplier ID is required.")

    if not invoice_number:
        fail("Invoice Number is required.")

    if not raw_material_code:
        fail("Raw Material Code is required.")

    if not finished_product:
        fail("Finished Product is required.")

    if not finished_product_code:
        fail("Finished Product Code is required.")

    input_weight = parse_decimal(input_weight_text, "Input Weight")
    output_weight = parse_decimal(output_weight_text, "Output Weight")
    initial_drc = parse_decimal(initial_drc_text, "Initial DRC")
    actual_drc = parse_decimal(actual_drc_text, "Actual DRC")

    if input_weight <= 0:
        fail("Input Weight must be greater than zero.")

    if output_weight < 0:
        fail("Output Weight cannot be negative.")

    if output_weight > input_weight:
        fail("Output Weight cannot exceed Input Weight.")

    if initial_drc < 0 or initial_drc > 100:
        fail("Initial DRC must be between 0 and 100.")

    if actual_drc < 0 or actual_drc > 100:
        fail("Actual DRC must be between 0 and 100.")

    database_path = get_database_path()

    try:
        workbook = load_workbook(database_path, data_only=False)
    except PermissionError:
        fail("Close Excel workbook before continuing.")
    except Exception as error:
        fail(f"Failed to open workbook.\n{error}")

    try:
        if PRODUCTION_SHEET_NAME not in workbook.sheetnames:
            fail(f"Sheet not found: {PRODUCTION_SHEET_NAME}")

        sheet = workbook[PRODUCTION_SHEET_NAME]
        quantity_available = calculate_available_quantity(
            workbook,
            supplier_id,
            invoice_number,
            raw_material_code,
        )

        if input_weight > quantity_available:
            fail(
                f"Input Weight exceeds available quantity. "
                f"Available quantity is {quantity_available:0.##}."
            )

        target_row = get_next_production_row(sheet)
        template_row = target_row - 1 if target_row > START_ROW else START_ROW
        batch_id = generate_next_batch_id(sheet)

        copy_row_style_and_formulas(sheet, template_row, target_row)

        sheet[f"{COL_SERIAL_NO}{target_row}"] = target_row - START_ROW + 1
        sheet[f"{COL_BATCH_ID}{target_row}"] = batch_id
        sheet[f"{COL_PRODUCTION_DATE}{target_row}"] = production_date
        sheet[f"{COL_PRODUCTION_DATE}{target_row}"].number_format = "dd-mm-yyyy"
        sheet[f"{COL_SUPPLIER_NAME}{target_row}"] = supplier_name
        sheet[f"{COL_SUPPLIER_ID}{target_row}"] = supplier_id
        sheet[f"{COL_INVOICE_NUMBER}{target_row}"] = invoice_number
        sheet[f"{COL_RAW_MATERIAL}{target_row}"] = raw_material
        sheet[f"{COL_RAW_MATERIAL_CODE}{target_row}"] = raw_material_code
        sheet[f"{COL_FINISHED_PRODUCT}{target_row}"] = finished_product
        sheet[f"{COL_FINISHED_PRODUCT_CODE}{target_row}"] = finished_product_code
        sheet[f"{COL_QUANTITY_AVAILABLE}{target_row}"] = quantity_available
        sheet[f"{COL_INPUT_WEIGHT}{target_row}"] = input_weight
        sheet[f"{COL_OUTPUT_WEIGHT}{target_row}"] = output_weight
        sheet[f"{COL_INITIAL_DRC}{target_row}"] = initial_drc
        sheet[f"{COL_ACTUAL_DRC}{target_row}"] = actual_drc

        write_calculation_formulas(sheet, target_row)

        workbook.calculation.fullCalcOnLoad = True
        workbook.calculation.forceFullCalc = True
        workbook.save(database_path)

        return f"Production batch {batch_id} saved successfully at row {target_row}."

    except PermissionError:
        fail("Cannot save workbook. Close Excel file first.")
    finally:
        workbook.close()


# ============================================================
# MAIN
# ============================================================

try:
    print(create_production_entry(sys.argv[1:]))
except Exception as error:
    fail(str(error))

