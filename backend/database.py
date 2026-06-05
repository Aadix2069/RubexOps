import json
from datetime import date, datetime
from pathlib import Path
from typing import Any, Dict, List, Optional

from openpyxl import load_workbook


# =========================================================
# CONFIG
# =========================================================

CONFIG_FOLDER = Path.home() / "Documents" / "RubexOps"
CONFIG_FILE = CONFIG_FOLDER / "database_config.json"

PCON_SHEET_NAME = "pcon"
PURCHASE_SHEET_NAME = "purchase"
SCON_SHEET_NAME = "scon"
SALES_SHEET_NAME = "sales"
PRODUCTION_SHEET_CANDIDATES = ("production", "prod", "production_data")

START_ROW = 5

# pcon columns (A:M)
PCON_USED_COLUMNS = tuple("BCDEFGHIJKLM")
PCON_REQUIRED_COLUMNS = ("B", "C", "D")

# purchase columns (A:N) - net_weight dynamically calculated, not stored
PURCHASE_USED_COLUMNS = tuple("BCDEFGHIJKLMN")
PURCHASE_REQUIRED_COLUMNS = ("B", "C", "D", "F")

# sales contract columns (A:M)
SCON_USED_COLUMNS = tuple("BCDEFGHIJKLM")
SCON_REQUIRED_COLUMNS = ("B", "C", "D")

# sales data columns (A:J)
SALES_USED_COLUMNS = tuple("BCDEFGHIJ")
SALES_REQUIRED_COLUMNS = ("B", "C", "D", "E", "F", "G")

# production columns (A:E) - Strictly input only
PRODUCTION_USED_COLUMNS = tuple("BCDE")
PRODUCTION_REQUIRED_COLUMNS = ("B", "C", "D", "E")


# =========================================================
# SAFE HELPERS
# =========================================================

def safe_string(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def safe_number(value: Any, default: float = 0.0) -> float:
    try:
        text = safe_string(value)
        if text == "":
            return float(default)
        return float(text.replace(",", ""))
    except Exception:
        return float(default)


def safe_int(value: Any, default: int = 0) -> int:
    try:
        text = safe_string(value)
        if text == "":
            return int(default)
        return int(float(text.replace(",", "")))
    except Exception:
        return int(default)


def safe_bool(value: Any, default: bool = False) -> bool:
    if value is None:
        return default

    if isinstance(value, bool):
        return value

    if isinstance(value, (int, float)):
        return value != 0

    text = safe_string(value).casefold()

    if text in {"1", "true", "yes", "y", "on"}:
        return True

    if text in {"0", "false", "no", "n", "off", ""}:
        return False

    return default


def safe_date(value: Any, default: str = "") -> str:
    if value is None:
        return default

    if isinstance(value, datetime):
        return value.strftime("%d-%m-%Y")

    if isinstance(value, date):
        return value.strftime("%d-%m-%Y")

    return safe_string(value)


def safe_cell(sheet, row: int, col: str) -> Any:
    return sheet[f"{col}{row}"].value


def row_has_data(sheet, row: int, columns: tuple[str, ...]) -> bool:
    for col in columns:
        value = safe_cell(sheet, row, col)
        if value not in (None, ""):
            return True
    return False


def row_is_empty(sheet, row: int, columns: tuple[str, ...]) -> bool:
    return not row_has_data(sheet, row, columns)


def next_serial_no(sheet, used_columns: tuple[str, ...], serial_col: str = "A") -> int:
    highest = 0

    for row in range(START_ROW, sheet.max_row + 1):
        if row_is_empty(sheet, row, used_columns):
            continue

        value = safe_cell(sheet, row, serial_col)
        if value is None:
            continue

        try:
            serial = int(float(str(value).strip()))
            if serial > highest:
                highest = serial
        except Exception:
            continue

    return highest + 1


def first_available_row(sheet, used_columns: tuple[str, ...]) -> int:
    for row in range(START_ROW, sheet.max_row + 1):
        if row_is_empty(sheet, row, used_columns):
            return row

    return sheet.max_row + 1 if sheet.max_row >= START_ROW else START_ROW


def _pick(data: Dict[str, Any], *keys: str, default: Any = "") -> Any:
    for key in keys:
        value = data.get(key)
        if value not in (None, ""):
            return value
    return default


def _get_sheet_by_candidates(workbook, candidates: tuple[str, ...], label: str):
    for name in candidates:
        if name in workbook.sheetnames:
            return workbook[name]

    raise ValueError(
        f"Sheet not found: {label}\nTried: {', '.join(candidates)}"
    )


# =========================================================
# DATABASE PATH
# =========================================================

def get_database_path() -> str:
    CONFIG_FOLDER.mkdir(parents=True, exist_ok=True)

    if not CONFIG_FILE.exists():
        raise FileNotFoundError(
            f"Database configuration not found:\n{CONFIG_FILE}"
        )

    with open(CONFIG_FILE, "r", encoding="utf-8") as f:
        config = json.load(f)

    database_path = safe_string(config.get("database_path"))

    if not database_path:
        raise ValueError("database_path missing in database_config.json")

    return database_path


# =========================================================
# WORKBOOK HELPERS
# =========================================================

def load_database():
    db_path = get_database_path()

    if not Path(db_path).exists():
        raise FileNotFoundError(
            f"Database file not found:\n{db_path}"
        )

    return load_workbook(db_path)


def save_database(workbook) -> None:
    db_path = get_database_path()
    workbook.save(db_path)


# =========================================================
# READ PCON
# =========================================================

def read_pcon() -> List[Dict[str, Any]]:
    wb = load_database()
    try:
        if PCON_SHEET_NAME not in wb.sheetnames:
            raise ValueError(f"Sheet not found: {PCON_SHEET_NAME}")

        sheet = wb[PCON_SHEET_NAME]
        contracts: List[Dict[str, Any]] = []

        for row in range(START_ROW, sheet.max_row + 1):
            if row_is_empty(sheet, row, PCON_REQUIRED_COLUMNS):
                continue

            contract_id = safe_string(safe_cell(sheet, row, "D"))
            if contract_id == "":
                continue

            contracts.append(
                {
                    "sl_no": safe_int(safe_cell(sheet, row, "A")),
                    "vendor_name": safe_string(safe_cell(sheet, row, "B")),
                    "vendor_id": safe_string(safe_cell(sheet, row, "C")),
                    "contract_id": contract_id,
                    "start_date": safe_date(safe_cell(sheet, row, "E")),
                    "end_date": safe_date(safe_cell(sheet, row, "F")),
                    "base_price": safe_number(safe_cell(sheet, row, "G")),
                    "agreed_qty": safe_number(safe_cell(sheet, row, "H")),
                    "breach_responsibility": safe_string(safe_cell(sheet, row, "I")),
                    "penalty_discount_percent": safe_number(safe_cell(sheet, row, "J")),
                    "remedy_days": safe_number(safe_cell(sheet, row, "K")),
                    "renewal_reference": safe_string(safe_cell(sheet, row, "L")),
                    "ignore_remaining_qty": safe_bool(safe_cell(sheet, row, "M")),
                }
            )

        return contracts
    finally:
        wb.close()


# =========================================================
# READ PURCHASE
# =========================================================

def read_purchase() -> List[Dict[str, Any]]:
    wb = load_database()
    try:
        if PURCHASE_SHEET_NAME not in wb.sheetnames:
            raise ValueError(f"Sheet not found: {PURCHASE_SHEET_NAME}")

        sheet = wb[PURCHASE_SHEET_NAME]
        purchases: List[Dict[str, Any]] = []

        for row in range(START_ROW, sheet.max_row + 1):
            if row_is_empty(sheet, row, PURCHASE_REQUIRED_COLUMNS):
                continue

            purchases.append(
                {
                    "sl_no": safe_int(safe_cell(sheet, row, "A")),
                    "vendor_id": safe_string(safe_cell(sheet, row, "B")),
                    "contract_id": safe_string(safe_cell(sheet, row, "C")),
                    "invoice_number": safe_string(safe_cell(sheet, row, "D")),
                    "purchase_order_date": safe_date(safe_cell(sheet, row, "E")),
                    "delivery_date": safe_date(safe_cell(sheet, row, "F")),
                    "invoice_weight": safe_number(safe_cell(sheet, row, "G")),
                    "before_unloading": safe_number(safe_cell(sheet, row, "H")),
                    "carrier_weight": safe_number(safe_cell(sheet, row, "I")),
                    "number_of_bags": safe_number(safe_cell(sheet, row, "J")),
                    "calculated_drc_percent": safe_number(safe_cell(sheet, row, "K")),
                    "gst_percent": safe_number(safe_cell(sheet, row, "L")),
                    "tds_percent": safe_number(safe_cell(sheet, row, "M")),
                    "unloading_charge": safe_number(safe_cell(sheet, row, "N")),
                }
            )

        return purchases
    finally:
        wb.close()


# =========================================================
# READ SCON
# =========================================================

def read_scon() -> List[Dict[str, Any]]:
    wb = load_database()
    try:
        if SCON_SHEET_NAME not in wb.sheetnames:
            raise ValueError(f"Sheet not found: {SCON_SHEET_NAME}")

        sheet = wb[SCON_SHEET_NAME]
        contracts: List[Dict[str, Any]] = []

        for row in range(START_ROW, sheet.max_row + 1):
            if row_is_empty(sheet, row, SCON_REQUIRED_COLUMNS):
                continue

            contract_id = safe_string(safe_cell(sheet, row, "D"))
            if contract_id == "":
                continue

            contracts.append(
                {
                    "sl_no": safe_int(safe_cell(sheet, row, "A")),
                    "customer_name": safe_string(safe_cell(sheet, row, "B")),
                    "customer_id": safe_string(safe_cell(sheet, row, "C")),
                    "contract_id": contract_id,
                    "start_date": safe_date(safe_cell(sheet, row, "E")),
                    "end_date": safe_date(safe_cell(sheet, row, "F")),
                    "base_price": safe_number(safe_cell(sheet, row, "G")),
                    "agreed_qty": safe_number(safe_cell(sheet, row, "H")),
                    "breach_responsibility": safe_string(safe_cell(sheet, row, "I")),
                    "penalty_discount_percent": safe_number(safe_cell(sheet, row, "J")),
                    "remedy_days": safe_number(safe_cell(sheet, row, "K")),
                    "renewal_reference": safe_string(safe_cell(sheet, row, "L")),
                    "ignore_remaining_qty": safe_bool(safe_cell(sheet, row, "M")),
                }
            )

        return contracts
    finally:
        wb.close()


# =========================================================
# READ SALES
# =========================================================

def read_sales() -> List[Dict[str, Any]]:
    wb = load_database()
    try:
        if SALES_SHEET_NAME not in wb.sheetnames:
            raise ValueError(f"Sheet not found: {SALES_SHEET_NAME}")

        sheet = wb[SALES_SHEET_NAME]
        sales_data: List[Dict[str, Any]] = []

        for row in range(START_ROW, sheet.max_row + 1):
            if row_is_empty(sheet, row, SALES_REQUIRED_COLUMNS):
                continue

            sales_data.append(
                {
                    "sl_no": safe_int(safe_cell(sheet, row, "A")),
                    "customer_id": safe_string(safe_cell(sheet, row, "B")),
                    "contract_id": safe_string(safe_cell(sheet, row, "C")),
                    "invoice_number": safe_string(safe_cell(sheet, row, "D")),
                    "sales_order_date": safe_date(safe_cell(sheet, row, "E")),
                    "dispatch_date": safe_date(safe_cell(sheet, row, "F")),
                    "weight": safe_number(safe_cell(sheet, row, "G")),
                    "gst_percent": safe_number(safe_cell(sheet, row, "H")),
                    "tcs_percent": safe_number(safe_cell(sheet, row, "I")),
                    "loading_charge": safe_number(safe_cell(sheet, row, "J")),
                }
            )

        return sales_data
    finally:
        wb.close()


# =========================================================
# READ PRODUCTION (UPDATED WITH ROW NUMBER)
# =========================================================

def read_production() -> List[Dict[str, Any]]:
    wb = load_database()
    try:
        sheet = _get_sheet_by_candidates(
            wb,
            PRODUCTION_SHEET_CANDIDATES,
            "production",
        )

        records: List[Dict[str, Any]] = []

        for row in range(START_ROW, sheet.max_row + 1):
            if row_is_empty(sheet, row, PRODUCTION_REQUIRED_COLUMNS):
                continue

            records.append(
                {
                    "row_number": row,  # <-- Added the exact row tracking identifier here
                    "sl_no": safe_int(safe_cell(sheet, row, "A")),
                    "batch_id": safe_string(safe_cell(sheet, row, "B")),
                    "production_date": safe_date(safe_cell(sheet, row, "C")),
                    "invoice_number": safe_string(safe_cell(sheet, row, "D")),
                    "output": safe_number(safe_cell(sheet, row, "E")),
                }
            )

        return records
    finally:
        wb.close()


# =========================================================
# APPEND PCON
# =========================================================

def append_pcon(contract_data: Dict[str, Any]) -> int:
    wb = load_database()
    try:
        if PCON_SHEET_NAME not in wb.sheetnames:
            raise ValueError(f"Sheet not found: {PCON_SHEET_NAME}")

        sheet = wb[PCON_SHEET_NAME]

        row = first_available_row(sheet, PCON_USED_COLUMNS)
        sl_no = next_serial_no(sheet, PCON_USED_COLUMNS)

        sheet[f"A{row}"] = sl_no
        sheet[f"B{row}"] = contract_data.get("vendor_name", "")
        sheet[f"C{row}"] = contract_data.get("vendor_id", "")
        sheet[f"D{row}"] = contract_data.get("contract_id", "")
        sheet[f"E{row}"] = contract_data.get("start_date", "")
        sheet[f"F{row}"] = contract_data.get("end_date", "")
        sheet[f"G{row}"] = contract_data.get("base_price", 0)
        sheet[f"H{row}"] = contract_data.get("agreed_qty", 0)
        sheet[f"I{row}"] = contract_data.get("breach_responsibility", "")
        sheet[f"J{row}"] = contract_data.get("penalty_discount_percent", 0)
        sheet[f"K{row}"] = contract_data.get("remedy_days", 0)
        sheet[f"L{row}"] = contract_data.get("renewal_reference", "")
        sheet[f"M{row}"] = 1 if safe_bool(contract_data.get("ignore_remaining_qty", False)) else 0

        save_database(wb)
        return sl_no
    finally:
        wb.close()


# =========================================================
# APPEND PURCHASE
# =========================================================

def append_purchase(purchase_data: Dict[str, Any]) -> int:
    wb = load_database()
    try:
        if PURCHASE_SHEET_NAME not in wb.sheetnames:
            raise ValueError(f"Sheet not found: {PURCHASE_SHEET_NAME}")

        sheet = wb[PURCHASE_SHEET_NAME]

        row = first_available_row(sheet, PURCHASE_USED_COLUMNS)
        sl_no = next_serial_no(sheet, PURCHASE_USED_COLUMNS)

        sheet[f"A{row}"] = sl_no
        sheet[f"B{row}"] = purchase_data.get("vendor_id", "")
        sheet[f"C{row}"] = purchase_data.get("contract_id", "")
        sheet[f"D{row}"] = purchase_data.get("invoice_number", "")
        sheet[f"E{row}"] = purchase_data.get("purchase_order_date", "")
        sheet[f"F{row}"] = purchase_data.get("delivery_date", "")
        sheet[f"G{row}"] = purchase_data.get("invoice_weight", 0)
        sheet[f"H{row}"] = purchase_data.get("before_unloading", 0)
        sheet[f"I{row}"] = purchase_data.get("carrier_weight", 0)
        sheet[f"J{row}"] = purchase_data.get("number_of_bags", 0)
        sheet[f"K{row}"] = purchase_data.get("calculated_drc_percent", 0)
        sheet[f"L{row}"] = purchase_data.get("gst_percent", 0)
        sheet[f"M{row}"] = purchase_data.get("tds_percent", 0)
        sheet[f"N{row}"] = purchase_data.get("unloading_charge", 0)

        save_database(wb)
        return sl_no
    finally:
        wb.close()


# =========================================================
# APPEND SCON
# =========================================================

def append_scon(contract_data: Dict[str, Any]) -> int:
    wb = load_database()
    try:
        if SCON_SHEET_NAME not in wb.sheetnames:
            raise ValueError(f"Sheet not found: {SCON_SHEET_NAME}")

        sheet = wb[SCON_SHEET_NAME]

        row = first_available_row(sheet, SCON_USED_COLUMNS)
        sl_no = next_serial_no(sheet, SCON_USED_COLUMNS)

        sheet[f"A{row}"] = sl_no
        sheet[f"B{row}"] = contract_data.get("customer_name", "")
        sheet[f"C{row}"] = contract_data.get("customer_id", "")
        sheet[f"D{row}"] = contract_data.get("contract_id", "")
        sheet[f"E{row}"] = contract_data.get("start_date", "")
        sheet[f"F{row}"] = contract_data.get("end_date", "")
        sheet[f"G{row}"] = contract_data.get("base_price", 0)
        sheet[f"H{row}"] = contract_data.get("agreed_qty", 0)
        sheet[f"I{row}"] = contract_data.get("breach_responsibility", "")
        sheet[f"J{row}"] = contract_data.get("penalty_discount_percent", 0)
        sheet[f"K{row}"] = contract_data.get("remedy_days", 0)
        sheet[f"L{row}"] = contract_data.get("renewal_reference", "")
        sheet[f"M{row}"] = 1 if safe_bool(contract_data.get("ignore_remaining_qty", False)) else 0

        save_database(wb)
        return sl_no
    finally:
        wb.close()


# =========================================================
# APPEND SALES
# =========================================================

def append_sales(sales_data: Dict[str, Any]) -> int:
    wb = load_database()
    try:
        if SALES_SHEET_NAME not in wb.sheetnames:
            raise ValueError(f"Sheet not found: {SALES_SHEET_NAME}")

        sheet = wb[SALES_SHEET_NAME]

        row = first_available_row(sheet, SALES_USED_COLUMNS)
        sl_no = next_serial_no(sheet, SALES_USED_COLUMNS)

        sheet[f"A{row}"] = sl_no
        sheet[f"B{row}"] = sales_data.get("customer_id", "")
        sheet[f"C{row}"] = sales_data.get("contract_id", "")
        sheet[f"D{row}"] = sales_data.get("invoice_number", "")
        sheet[f"E{row}"] = sales_data.get("sales_order_date", "")
        sheet[f"F{row}"] = sales_data.get("dispatch_date", "")
        sheet[f"G{row}"] = sales_data.get("weight", 0)
        sheet[f"H{row}"] = sales_data.get("gst_percent", 0)
        sheet[f"I{row}"] = sales_data.get("tcs_percent", 0)
        sheet[f"J{row}"] = sales_data.get("loading_charge", 0)

        save_database(wb)
        return sl_no
    finally:
        wb.close()


# =========================================================
# PRODUCTION HELPERS
# =========================================================

def get_purchase_by_invoice(invoice_number: str) -> Optional[Dict[str, Any]]:
    invoice_key = safe_string(invoice_number).casefold()
    if invoice_key == "":
        return None

    for purchase in read_purchase():
        if safe_string(purchase.get("invoice_number")).casefold() == invoice_key:
            return purchase

    return None


def production_invoice_exists(invoice_number: str) -> bool:
    invoice_key = safe_string(invoice_number).casefold()
    if invoice_key == "":
        return False

    wb = load_database()
    try:
        sheet = _get_sheet_by_candidates(
            wb,
            PRODUCTION_SHEET_CANDIDATES,
            "production",
        )

        for row in range(START_ROW, sheet.max_row + 1):
            current_invoice = safe_string(safe_cell(sheet, row, "D")).casefold()
            if current_invoice == invoice_key:
                return True

        return False
    finally:
        wb.close()


def _parse_date_for_sort(date_str: str) -> date:
    if not date_str:
        return date.min
    try:
        return datetime.strptime(date_str, "%d-%m-%Y").date()
    except ValueError:
        return date.min


def generate_production_batch_id(invoice_number: str) -> str:
    """
    Finds the purchase, sorts all purchases for that contract chronologically,
    and returns a formatted Batch ID (e.g., '123617-R1-P012').
    """
    purchase = get_purchase_by_invoice(invoice_number)
    if not purchase:
        raise ValueError("Cannot generate Batch ID: Purchase invoice not found.")

    contract_id = safe_string(purchase.get("contract_id"))
    if not contract_id:
        raise ValueError("Cannot generate Batch ID: Invoice has no Contract ID.")

    all_purchases = read_purchase()
    contract_purchases = [
        p for p in all_purchases
        if safe_string(p.get("contract_id")).casefold() == contract_id.casefold()
    ]

    # Sort purchases chronologically by PO Date (and fallback to sl_no for ties)
    contract_purchases.sort(
        key=lambda p: (
            _parse_date_for_sort(safe_string(p.get("purchase_order_date"))),
            safe_int(p.get("sl_no"))
        )
    )

    target_invoice = safe_string(invoice_number).casefold()
    index = 0
    for i, p in enumerate(contract_purchases, start=1):
        if safe_string(p.get("invoice_number")).casefold() == target_invoice:
            index = i
            break

    # Safety fallback (should never happen if logic is intact)
    if index == 0:
        index = len(contract_purchases) + 1

    return f"{contract_id}-P{index:03d}"


def append_production(production_data: Dict[str, Any]) -> int:
    """
    Expects strictly the user-input parameters:
      - invoice_number
      - production_date
      - output
    """
    wb = load_database()
    try:
        sheet = _get_sheet_by_candidates(
            wb,
            PRODUCTION_SHEET_CANDIDATES,
            "production",
        )

        invoice_number = safe_string(production_data.get("invoice_number"))
        if not invoice_number:
            raise ValueError("Invoice Number is required.")

        if production_invoice_exists(invoice_number):
            raise ValueError("This purchase invoice has already been used for a production batch.")

        batch_id = generate_production_batch_id(invoice_number)

        production_date = _pick(production_data, "production_date")
        if not production_date:
            raise ValueError("Production Date is required.")

        output_qty = safe_number(production_data.get("output"))
        if output_qty <= 0:
            raise ValueError("Output must be greater than zero.")

        row = first_available_row(sheet, PRODUCTION_USED_COLUMNS)
        sl_no = next_serial_no(sheet, PRODUCTION_USED_COLUMNS)

        sheet[f"A{row}"] = sl_no
        sheet[f"B{row}"] = batch_id
        sheet[f"C{row}"] = safe_date(production_date) if isinstance(production_date, (date, datetime)) else safe_string(production_date)
        sheet[f"D{row}"] = invoice_number
        sheet[f"E{row}"] = output_qty

        save_database(wb)
        return sl_no
    finally:
        wb.close()


def read_production_by_invoice(invoice_number: str) -> Optional[Dict[str, Any]]:
    invoice_key = safe_string(invoice_number).casefold()
    if invoice_key == "":
        return None

    for record in read_production():
        if safe_string(record.get("invoice_number")).casefold() == invoice_key:
            return record

    return None


def read_production_by_batch_id(batch_id: str) -> Optional[Dict[str, Any]]:
    batch_key = safe_string(batch_id).casefold()
    if batch_key == "":
        return None

    for record in read_production():
        if safe_string(record.get("batch_id")).casefold() == batch_key:
            return record

    return None


# =========================================================
# OPTIONAL DELETE HELPERS (PCON)
# =========================================================

def delete_pcon_by_contract_id(contract_id: str) -> int:
    wb = load_database()
    try:
        if PCON_SHEET_NAME not in wb.sheetnames:
            raise ValueError(f"Sheet not found: {PCON_SHEET_NAME}")

        sheet = wb[PCON_SHEET_NAME]
        contract_id = safe_string(contract_id)

        deleted = 0
        for row in range(sheet.max_row, START_ROW - 1, -1):
            if safe_string(safe_cell(sheet, row, "D")) == contract_id:
                sheet.delete_rows(row, 1)
                deleted += 1

        if deleted:
            save_database(wb)

        return deleted
    finally:
        wb.close()


def delete_purchase_by_contract_id(contract_id: str) -> int:
    wb = load_database()
    try:
        if PURCHASE_SHEET_NAME not in wb.sheetnames:
            raise ValueError(f"Sheet not found: {PURCHASE_SHEET_NAME}")

        sheet = wb[PURCHASE_SHEET_NAME]
        contract_id = safe_string(contract_id)

        deleted = 0
        for row in range(sheet.max_row, START_ROW - 1, -1):
            if safe_string(safe_cell(sheet, row, "C")) == contract_id:
                sheet.delete_rows(row, 1)
                deleted += 1

        if deleted:
            save_database(wb)

        return deleted
    finally:
        wb.close()


# =========================================================
# OPTIONAL DELETE HELPERS (SCON / SALES)
# =========================================================

def delete_scon_by_contract_id(contract_id: str) -> int:
    wb = load_database()
    try:
        if SCON_SHEET_NAME not in wb.sheetnames:
            raise ValueError(f"Sheet not found: {SCON_SHEET_NAME}")

        sheet = wb[SCON_SHEET_NAME]
        contract_id = safe_string(contract_id)

        deleted = 0
        for row in range(sheet.max_row, START_ROW - 1, -1):
            if safe_string(safe_cell(sheet, row, "D")) == contract_id:
                sheet.delete_rows(row, 1)
                deleted += 1

        if deleted:
            save_database(wb)

        return deleted
    finally:
        wb.close()


def delete_sales_by_contract_id(contract_id: str) -> int:
    wb = load_database()
    try:
        if SALES_SHEET_NAME not in wb.sheetnames:
            raise ValueError(f"Sheet not found: {SALES_SHEET_NAME}")

        sheet = wb[SALES_SHEET_NAME]
        contract_id = safe_string(contract_id)

        deleted = 0
        for row in range(sheet.max_row, START_ROW - 1, -1):
            if safe_string(safe_cell(sheet, row, "C")) == contract_id:
                sheet.delete_rows(row, 1)
                deleted += 1

        if deleted:
            save_database(wb)

        return deleted
    finally:
        wb.close()