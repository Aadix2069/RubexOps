import json
import os
import sys
from copy import copy
from datetime import datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.formula.translate import Translator


# ============================================================
# CONFIG
# ============================================================

CONFIG_PATH = (
    Path.home()
    / "Documents"
    / "RubexOps"
    / "database_config.json"
)

SALES_SHEET_NAME = "sales"
SCON_SHEET_NAME = "scon"
START_ROW = 5


# ============================================================
# COLUMN MAP
# ============================================================

COL_SERIAL_NO = "A"
COL_CUSTOMER_ID = "C"
COL_ITEM_CODE = "E"
COL_INVOICE_NUMBER = "F"
COL_SALES_ORDER = "G"
COL_DISPATCH = "H"
COL_WEIGHT = "I"
COL_BASE_RATE = "J"
COL_GST_PERCENT = "K"
COL_TCS_PERCENT = "L"
COL_LOADING_CHARGE = "P"


# ============================================================
# ERROR HANDLER
# ============================================================

def fail(message):
    print(f"ERROR: {message}", file=sys.stderr)
    sys.exit(1)


# ============================================================
# HELPERS
# ============================================================

def safe_string(value):
    if value is None:
        return ""

    return str(value).strip()


def get_database_path():
    if not CONFIG_PATH.exists():
        fail(f"Database config file not found: {CONFIG_PATH}")

    try:
        with open(CONFIG_PATH, "r", encoding="utf-8") as file:
            config = json.load(file)
    except json.JSONDecodeError:
        fail("Invalid JSON inside database_config.json.")

    database_path = (
        config.get("database_path")
        or config.get("excel_path")
        or config.get("path")
        or ""
    ).strip()

    if not database_path:
        fail("Database path was not found in database_config.json.")

    if not os.path.exists(database_path):
        fail(f"Excel database file not found: {database_path}")

    return database_path


def parse_date(value, field_name):
    text = safe_string(value)

    if not text:
        fail(f"{field_name} is required.")

    try:
        return datetime.strptime(text, "%d-%m-%Y")
    except ValueError:
        fail(f"{field_name} must be in dd-MM-yyyy format.")


def parse_decimal(value, field_name, required=True):
    text = safe_string(value).replace(",", "").replace("%", "")

    if not text:
        if required:
            fail(f"{field_name} is required.")

        return None

    try:
        return float(Decimal(text))
    except (InvalidOperation, ValueError):
        fail(f"{field_name} must be numeric.")


def get_next_sales_row(sheet):
    row = START_ROW

    while True:
        customer_id = sheet[f"{COL_CUSTOMER_ID}{row}"].value
        invoice_number = sheet[f"{COL_INVOICE_NUMBER}{row}"].value

        if safe_string(customer_id) == "" and safe_string(invoice_number) == "":
            return row

        row += 1


def copy_row_style_and_formulas(sheet, source_row, target_row):
    for column in range(1, sheet.max_column + 1):
        source_cell = sheet.cell(row=source_row, column=column)
        target_cell = sheet.cell(row=target_row, column=column)

        if source_cell.has_style:
            target_cell.font = copy(source_cell.font)
            target_cell.fill = copy(source_cell.fill)
            target_cell.border = copy(source_cell.border)
            target_cell.alignment = copy(source_cell.alignment)
            target_cell.number_format = source_cell.number_format
            target_cell.protection = copy(source_cell.protection)

        if isinstance(source_cell.value, str) and source_cell.value.startswith("="):
            target_cell.value = Translator(
                source_cell.value,
                origin=source_cell.coordinate
            ).translate_formula(target_cell.coordinate)

    if source_row in sheet.row_dimensions:
        sheet.row_dimensions[target_row].height = sheet.row_dimensions[source_row].height


def ensure_customer_contract_exists(workbook, customer_id, item_code):
    if SCON_SHEET_NAME not in workbook.sheetnames:
        fail(f"Sheet not found: {SCON_SHEET_NAME}")

    sheet = workbook[SCON_SHEET_NAME]
    customer_key = customer_id.casefold()
    item_key = item_code.casefold()

    for row in range(START_ROW, sheet.max_row + 1):
        existing_customer_id = safe_string(sheet[f"C{row}"].value)
        existing_item_code = safe_string(sheet[f"E{row}"].value)

        if existing_customer_id.casefold() == customer_key and existing_item_code.casefold() == item_key:
            return

    fail("Selected customer and item were not found in the sales contract sheet.")


def find_duplicate_invoice_row(sheet, customer_id, invoice_number):
    customer_key = customer_id.casefold()
    invoice_key = invoice_number.casefold()

    for row in range(START_ROW, sheet.max_row + 1):
        existing_customer_id = safe_string(sheet[f"{COL_CUSTOMER_ID}{row}"].value)
        existing_invoice_number = safe_string(sheet[f"{COL_INVOICE_NUMBER}{row}"].value)

        if existing_customer_id.casefold() == customer_key and existing_invoice_number.casefold() == invoice_key:
            return row

    return None


# ============================================================
# CREATE SALES DATA
# ============================================================

def create_sales_data(args):
    if len(args) != 10:
        fail(
            "Expected 10 arguments: customer id, item code, invoice number, "
            "sales order date, dispatch date, weight, base rate, GST, "
            "TCS 194Q, loading charge."
        )

    (
        customer_id,
        item_code,
        invoice_number,
        sales_order_date_text,
        dispatch_date_text,
        weight_text,
        base_rate_text,
        gst_percent_text,
        tcs_percent_text,
        loading_charge_text,
    ) = args

    customer_id = safe_string(customer_id)
    item_code = safe_string(item_code)
    invoice_number = safe_string(invoice_number)

    if not customer_id:
        fail("Customer ID is required.")

    if not item_code:
        fail("Item Code is required.")

    if not invoice_number:
        fail("Invoice Number is required.")

    sales_order_date = parse_date(sales_order_date_text, "Sales Order date")
    dispatch_date = parse_date(dispatch_date_text, "Dispatch date")

    if sales_order_date > dispatch_date:
        fail("Dispatch date must be after Sales Order date.")

    weight = parse_decimal(weight_text, "Weight")
    base_rate = parse_decimal(base_rate_text, "Base Rate")
    gst_percent = parse_decimal(gst_percent_text, "GST")
    tcs_percent = parse_decimal(tcs_percent_text, "TCS 194Q")
    loading_charge = parse_decimal(loading_charge_text, "Loading Charge", required=False)

    if weight <= 0:
        fail("Weight must be greater than zero.")

    if base_rate <= 0:
        fail("Base Rate must be greater than zero.")

    if gst_percent < 0 or gst_percent > 100:
        fail("GST must be between 0 and 100.")

    if tcs_percent < 0 or tcs_percent > 100:
        fail("TCS 194Q must be between 0 and 100.")

    if loading_charge is not None and loading_charge < 0:
        fail("Loading Charge cannot be negative.")

    database_path = get_database_path()

    try:
        workbook = load_workbook(database_path, data_only=False)
    except PermissionError:
        fail("Excel database is open or locked. Please close the workbook and try again.")
    except Exception as error:
        fail(f"Failed to open workbook.\n{error}")

    try:
        if SALES_SHEET_NAME not in workbook.sheetnames:
            fail(f"Sheet not found: {SALES_SHEET_NAME}")

        ensure_customer_contract_exists(workbook, customer_id, item_code)

        sheet = workbook[SALES_SHEET_NAME]
        duplicate_row = find_duplicate_invoice_row(sheet, customer_id, invoice_number)

        if duplicate_row is not None:
            fail(
                "This invoice number already exists for the selected customer "
                f"at sales sheet row {duplicate_row}."
            )

        target_row = get_next_sales_row(sheet)
        template_row = target_row - 1 if target_row > START_ROW else START_ROW

        copy_row_style_and_formulas(sheet, template_row, target_row)

        sheet[f"{COL_SERIAL_NO}{target_row}"] = target_row - START_ROW + 1
        sheet[f"{COL_CUSTOMER_ID}{target_row}"] = customer_id
        sheet[f"{COL_ITEM_CODE}{target_row}"] = item_code
        sheet[f"{COL_INVOICE_NUMBER}{target_row}"] = invoice_number
        sheet[f"{COL_SALES_ORDER}{target_row}"] = sales_order_date
        sheet[f"{COL_DISPATCH}{target_row}"] = dispatch_date
        sheet[f"{COL_SALES_ORDER}{target_row}"].number_format = "dd-mm-yyyy"
        sheet[f"{COL_DISPATCH}{target_row}"].number_format = "dd-mm-yyyy"
        sheet[f"{COL_WEIGHT}{target_row}"] = weight
        sheet[f"{COL_BASE_RATE}{target_row}"] = base_rate
        sheet[f"{COL_GST_PERCENT}{target_row}"] = gst_percent
        sheet[f"{COL_TCS_PERCENT}{target_row}"] = tcs_percent
        sheet[f"{COL_LOADING_CHARGE}{target_row}"] = loading_charge

        workbook.calculation.fullCalcOnLoad = True
        workbook.calculation.forceFullCalc = True
        workbook.save(database_path)

        return f"Sales data saved successfully at row {target_row}."

    except PermissionError:
        fail("Cannot save workbook. Close Excel file first.")
    finally:
        workbook.close()


try:
    print(create_sales_data(sys.argv[1:]))
except Exception as error:
    fail(str(error))

