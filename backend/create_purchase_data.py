import json
import os
import sys
from copy import copy
from datetime import datetime, date
from pathlib import Path

from openpyxl.formula.translate import Translator
from openpyxl import load_workbook

# ============================================================
# CONFIG
# ============================================================

CONFIG_PATH = (
    Path.home()
    / "Documents"
    / "RubexOps"
    / "database_config.json"
)

PCON_SHEET_NAME     = "pcon"
PURCHASE_SHEET_NAME = "purchase"
START_ROW           = 5

# ============================================================
# COLUMN MAP  (purchase sheet)
# ============================================================

COL_SERIAL_NO       = "A"
COL_VENDOR_ID       = "C"
COL_INVOICE_NUMBER  = "F"
COL_PURCHASE_ORDER  = "G"
COL_DELIVERY        = "H"
COL_INVOICE_WEIGHT  = "I"
COL_BEFORE_UNLOAD   = "J"
COL_CARRIER_WEIGHT  = "K"
COL_NO_OF_BAGS      = "M"
COL_CALCULATED_DRC  = "O"
COL_GST             = "S"   
COL_TDS             = "T"   
COL_UNLOADING       = "V"   

# ============================================================
# DATABASE PATH
# ============================================================

def get_database_path():
    if not CONFIG_PATH.exists():
        raise FileNotFoundError(f"Database config file not found: {CONFIG_PATH}")
    with open(CONFIG_PATH, "r", encoding="utf-8") as file:
        config = json.load(file)
    database_path = config.get("database_path") or config.get("excel_path") or config.get("path")
    if not database_path:
        raise ValueError("Database path was not found in database_config.json.")
    if not os.path.exists(database_path):
        raise FileNotFoundError(f"Excel database file not found: {database_path}")
    return database_path

# ============================================================
# CONTRACT VALIDATION
# ============================================================

def verify_contract_active(sheet, vendor_id, purchase_order_date):
    today = date.today()
    found = False

    for row in range(START_ROW, sheet.max_row + 1):
        cid = sheet[f"C{row}"].value
        if cid is None:
            continue
        if str(cid).strip() != str(vendor_id).strip():
            continue

        found = True
        start_raw      = sheet[f"F{row}"].value
        end_raw        = sheet[f"G{row}"].value
        
        remaining_qty  = sheet[f"M{row}"].value 
        breach         = sheet[f"O{row}"].value 

        if hasattr(start_raw, "date"):
            contract_start = start_raw.date()
        else:
            try:
                contract_start = datetime.strptime(str(start_raw).strip(), "%d-%m-%Y").date()
            except Exception:
                raise ValueError(f"Cannot parse contract start date for Vendor ID {vendor_id}.")

        if hasattr(end_raw, "date"):
            contract_end = end_raw.date()
        else:
            try:
                contract_end = datetime.strptime(str(end_raw).strip(), "%d-%m-%Y").date()
            except Exception:
                raise ValueError(f"Cannot parse contract end date for Vendor ID {vendor_id}.")

        try:
            po_date = datetime.strptime(str(purchase_order_date).strip(), "%d-%m-%Y").date()
        except Exception:
            raise ValueError(f"Invalid Purchase Order Date format: {purchase_order_date}. Expected dd-MM-yyyy.")

        try:
            remaining = float(remaining_qty) if remaining_qty not in (None, "") else 0
        except (TypeError, ValueError):
            remaining = 0

        breach_yes = str(breach).strip().upper() == "YES" if breach else False

        if today > contract_end:
            raise ValueError(f"The contract for Vendor ID '{vendor_id}' has expired (End Date: {contract_end.strftime('%d-%m-%Y')}). Please renew the contract before entering purchases.")
        if today < contract_start:
            raise ValueError(f"The contract for Vendor ID '{vendor_id}' has not started yet (Start Date: {contract_start.strftime('%d-%m-%Y')}).")
        if remaining <= 0:
            raise ValueError(f"The contract for Vendor ID '{vendor_id}' is already completed (Remaining Quantity: 0). No further purchases can be entered.")
        if breach_yes:
            raise ValueError(f"The contract for Vendor ID '{vendor_id}' is in breach. Resolve the breach before entering new purchases.")

        if po_date < contract_start or po_date > contract_end:
            raise ValueError(f"Purchase Order Date ({po_date.strftime('%d-%m-%Y')}) is outside the contract period ({contract_start.strftime('%d-%m-%Y')} — {contract_end.strftime('%d-%m-%Y')}).")
        return

    if not found:
        raise ValueError(f"No active contract found for Vendor ID '{vendor_id}'. Please select a valid active contract.")

# ============================================================
# ROW HELPERS
# ============================================================

def get_next_purchase_row(sheet):
    row = START_ROW
    while True:
        vendor_id      = sheet[f"{COL_VENDOR_ID}{row}"].value
        invoice_number = sheet[f"{COL_INVOICE_NUMBER}{row}"].value
        if vendor_id is None and invoice_number is None:
            return row
        row += 1

def copy_row_style_and_formulas(sheet, source_row, target_row):
    for column in range(1, sheet.max_column + 1):
        source_cell = sheet.cell(row=source_row, column=column)
        target_cell = sheet.cell(row=target_row, column=column)
        if source_cell.has_style:
            target_cell.font          = copy(source_cell.font)
            target_cell.fill          = copy(source_cell.fill)
            target_cell.border        = copy(source_cell.border)
            target_cell.alignment     = copy(source_cell.alignment)
            target_cell.number_format = source_cell.number_format
            target_cell.protection    = copy(source_cell.protection)
        if isinstance(source_cell.value, str) and source_cell.value.startswith("="):
            target_cell.value = Translator(source_cell.value, origin=source_cell.coordinate).translate_formula(target_cell.coordinate)
    if source_row in sheet.row_dimensions:
        sheet.row_dimensions[target_row].height = sheet.row_dimensions[source_row].height

# ============================================================
# WRITE PURCHASE DATA
# ============================================================

def create_purchase_data(args):
    if len(args) != 12:
        raise ValueError("Expected 12 arguments.")

    (vendor_id, invoice_number, purchase_order_date, delivery_date, invoice_weight, 
     before_unloading, carrier_weight, no_of_bags, calculated_drc, gst, tds, unloading_charge) = args

    def to_float_or_none(value, field_name):
        if value == "" or value is None: return None
        try: result = float(value)
        except (ValueError, TypeError): raise ValueError(f"{field_name} must be numeric.")
        if result < 0: raise ValueError(f"{field_name} cannot be negative.")
        return result

    def to_int_or_none(value, field_name):
        if value == "" or value is None: return None
        try: result = int(value)
        except (ValueError, TypeError): raise ValueError(f"{field_name} must be a whole number.")
        if result < 0: raise ValueError(f"{field_name} cannot be negative.")
        return result

    def to_float_required(value, field_name):
        if value == "" or value is None: raise ValueError(f"{field_name} is required.")
        try: result = float(value)
        except (ValueError, TypeError): raise ValueError(f"{field_name} must be numeric.")
        if result < 0: raise ValueError(f"{field_name} cannot be negative.")
        return result

    before_unloading_val = to_float_required(before_unloading, "Before Unloading")
    carrier_weight_val   = to_float_required(carrier_weight,   "Carrier Weight")
    calculated_drc_val   = to_float_required(calculated_drc,   "Calculated DRC")
    
    # PERCENTAGE BUG FIX: Divide exactly by 100 to map to raw Excel float logic safely
    gst_val              = to_float_required(gst,              "GST") / 100.0
    tds_val              = to_float_required(tds,              "TDS 194Q") / 100.0

    if not (0 <= calculated_drc_val <= 100): raise ValueError("Calculated DRC must be between 0 and 100.")
    if gst_val < 0: raise ValueError("GST cannot be negative.")
    if tds_val < 0 or tds_val > 1: raise ValueError("TDS 194Q must be between 0 and 100.")

    invoice_weight_val   = to_float_or_none(invoice_weight,   "Invoice Weight")
    no_of_bags_val       = to_int_or_none(no_of_bags,         "No. of Bags")
    unloading_val        = to_float_or_none(unloading_charge, "Unloading Charge")

    delivery_date_val = None
    if delivery_date.strip():
        try:
            po  = datetime.strptime(purchase_order_date.strip(), "%d-%m-%Y")
            dlv = datetime.strptime(delivery_date.strip(),       "%d-%m-%Y")
        except ValueError as exc:
            raise ValueError(f"Invalid date format. Expected dd-MM-yyyy.")
        if dlv < po: raise ValueError("Delivery Date cannot be before the Purchase Order Date.")
        delivery_date_val = dlv

    po_date_val = datetime.strptime(purchase_order_date.strip(), "%d-%m-%Y")

    database_path = get_database_path()
    workbook = load_workbook(database_path, data_only=False)

    if PCON_SHEET_NAME not in workbook.sheetnames: raise ValueError(f"Sheet not found: {PCON_SHEET_NAME}")
    pcon_sheet = workbook[PCON_SHEET_NAME]
    verify_contract_active(pcon_sheet, vendor_id, purchase_order_date)

    if PURCHASE_SHEET_NAME not in workbook.sheetnames: raise ValueError(f"Sheet not found: {PURCHASE_SHEET_NAME}")
    sheet = workbook[PURCHASE_SHEET_NAME]

    target_row = get_next_purchase_row(sheet)
    template_row = target_row - 1 if target_row > START_ROW else START_ROW
    copy_row_style_and_formulas(sheet, template_row, target_row)

    sheet[f"{COL_SERIAL_NO}{target_row}"] = target_row - START_ROW + 1
    sheet[f"{COL_VENDOR_ID}{target_row}"]      = vendor_id
    sheet[f"{COL_INVOICE_NUMBER}{target_row}"] = invoice_number
    sheet[f"{COL_PURCHASE_ORDER}{target_row}"] = po_date_val

    if delivery_date_val is not None: sheet[f"{COL_DELIVERY}{target_row}"] = delivery_date_val
    if invoice_weight_val is not None: sheet[f"{COL_INVOICE_WEIGHT}{target_row}"] = invoice_weight_val

    sheet[f"{COL_BEFORE_UNLOAD}{target_row}"]  = before_unloading_val
    sheet[f"{COL_CARRIER_WEIGHT}{target_row}"] = carrier_weight_val

    if no_of_bags_val is not None: sheet[f"{COL_NO_OF_BAGS}{target_row}"] = no_of_bags_val

    sheet[f"{COL_CALCULATED_DRC}{target_row}"] = calculated_drc_val
    sheet[f"{COL_GST}{target_row}"]            = gst_val
    sheet[f"{COL_TDS}{target_row}"]            = tds_val

    if unloading_val is not None: sheet[f"{COL_UNLOADING}{target_row}"] = unloading_val

    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc  = True
    workbook.save(database_path)

    return f"Purchase entry saved successfully at row {target_row}.\nVendor ID: {vendor_id}  |  Invoice: {invoice_number}  |  PO Date: {purchase_order_date}"

# ============================================================
# MAIN
# ============================================================
try:
    print(create_purchase_data(sys.argv[1:]))
except Exception as error:
    print(f"ERROR: {error}", file=sys.stderr)
    sys.exit(1)