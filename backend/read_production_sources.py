import json
import os
import sys
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import load_workbook


# ============================================================
# CONFIG
# ============================================================

CONFIG_PATH = (
    Path.home()
    / "Documents"
    / "RubexOps"
    / "database_config.json"
)

PURCHASE_SHEET_NAME = "purchase"
PRODUCTION_SHEET_NAME = "production"
START_ROW = 5


# ============================================================
# PURCHASE COLUMN MAP
# ============================================================

PUR_COL_SUPPLIER_NAME = "B"
PUR_COL_SUPPLIER_ID = "C"
PUR_COL_RAW_MATERIAL = "D"
PUR_COL_RAW_MATERIAL_CODE = "E"
PUR_COL_INVOICE_NUMBER = "F"
PUR_COL_INVOICE_WEIGHT = "I"
PUR_COL_NET_WEIGHT = "N"
PUR_COL_DRC_WEIGHT = "P"


# ============================================================
# PRODUCTION COLUMN MAP
# ============================================================

PROD_COL_BATCH_ID = "B"
PROD_COL_SUPPLIER_ID = "E"
PROD_COL_INVOICE_NUMBER = "F"
PROD_COL_RAW_MATERIAL_CODE = "H"
PROD_COL_INPUT_WEIGHT = "L"


# ============================================================
# ERROR HANDLER
# ============================================================

def fail(message):
    print(f"ERROR: {message}", file=sys.stderr)
    sys.exit(1)


# ============================================================
# SAFE HELPERS
# ============================================================

def safe_string(value):
    if value is None:
        return ""

    return str(value).strip()


def number_or_zero(value):
    if value is None:
        return 0.0

    try:
        if isinstance(value, str):
            text = value.strip().replace(",", "").replace("%", "")

            if not text:
                return 0.0

            return float(Decimal(text))

        return float(value)
    except (InvalidOperation, ValueError, TypeError):
        return 0.0


def get_database_path():
    if not CONFIG_PATH.exists():
        fail(f"Database config file not found: {CONFIG_PATH}")

    try:
        with open(CONFIG_PATH, "r", encoding="utf-8") as file:
            config = json.load(file)
    except json.JSONDecodeError:
        fail("Invalid JSON inside database_config.json.")

    database_path = (
        config.get("database_path")
        or config.get("excel_path")
        or config.get("path")
        or ""
    ).strip()

    if not database_path:
        fail("Database path was not found in database_config.json.")

    if not os.path.exists(database_path):
        fail(f"Excel database file not found: {database_path}")

    return database_path


# ============================================================
# INVENTORY HELPERS
# ============================================================

def stock_key(supplier_id, invoice_number, raw_material_code):
    return (
        safe_string(supplier_id).casefold(),
        safe_string(invoice_number).casefold(),
        safe_string(raw_material_code).casefold(),
    )


def get_purchase_quantity(sheet, row):
    net_weight = number_or_zero(sheet[f"{PUR_COL_NET_WEIGHT}{row}"].value)

    if net_weight > 0:
        return net_weight

    drc_weight = number_or_zero(sheet[f"{PUR_COL_DRC_WEIGHT}{row}"].value)

    if drc_weight > 0:
        return drc_weight

    return number_or_zero(sheet[f"{PUR_COL_INVOICE_WEIGHT}{row}"].value)


def read_consumption(workbook):
    consumption = {}

    if PRODUCTION_SHEET_NAME not in workbook.sheetnames:
        return consumption

    sheet = workbook[PRODUCTION_SHEET_NAME]

    for row in range(START_ROW, sheet.max_row + 1):
        supplier_id = safe_string(sheet[f"{PROD_COL_SUPPLIER_ID}{row}"].value)
        invoice_number = safe_string(sheet[f"{PROD_COL_INVOICE_NUMBER}{row}"].value)
        raw_material_code = safe_string(sheet[f"{PROD_COL_RAW_MATERIAL_CODE}{row}"].value)

        if not supplier_id or not invoice_number or not raw_material_code:
            continue

        key = stock_key(supplier_id, invoice_number, raw_material_code)
        consumption[key] = consumption.get(key, 0.0) + number_or_zero(
            sheet[f"{PROD_COL_INPUT_WEIGHT}{row}"].value
        )

    return consumption


def generate_next_batch_id(workbook):
    highest = 0

    if PRODUCTION_SHEET_NAME not in workbook.sheetnames:
        return "BATCH-0001"

    sheet = workbook[PRODUCTION_SHEET_NAME]

    for row in range(START_ROW, sheet.max_row + 1):
        batch_id = safe_string(sheet[f"{PROD_COL_BATCH_ID}{row}"].value)

        if not batch_id.upper().startswith("BATCH-"):
            continue

        try:
            number = int(batch_id.split("-", 1)[1])
        except (IndexError, ValueError):
            continue

        highest = max(highest, number)

    return f"BATCH-{highest + 1:04d}"


# ============================================================
# READ SOURCES
# ============================================================

def read_sources():
    database_path = get_database_path()

    try:
        workbook = load_workbook(database_path, data_only=True, read_only=True)
    except PermissionError:
        fail("Close Excel workbook before continuing.")
    except Exception as error:
        fail(f"Failed to open workbook.\n{error}")

    try:
        if PURCHASE_SHEET_NAME not in workbook.sheetnames:
            fail(f"Sheet not found: {PURCHASE_SHEET_NAME}")

        purchase_sheet = workbook[PURCHASE_SHEET_NAME]
        consumption = read_consumption(workbook)
        source_map = {}

        for row in range(START_ROW, purchase_sheet.max_row + 1):
            supplier_name = safe_string(purchase_sheet[f"{PUR_COL_SUPPLIER_NAME}{row}"].value)
            supplier_id = safe_string(purchase_sheet[f"{PUR_COL_SUPPLIER_ID}{row}"].value)
            invoice_number = safe_string(purchase_sheet[f"{PUR_COL_INVOICE_NUMBER}{row}"].value)
            raw_material = safe_string(purchase_sheet[f"{PUR_COL_RAW_MATERIAL}{row}"].value)
            raw_material_code = safe_string(purchase_sheet[f"{PUR_COL_RAW_MATERIAL_CODE}{row}"].value)

            if not supplier_id or not invoice_number or not raw_material_code:
                continue

            quantity_purchased = get_purchase_quantity(purchase_sheet, row)

            if quantity_purchased <= 0:
                continue

            key = stock_key(supplier_id, invoice_number, raw_material_code)

            if key not in source_map:
                source_map[key] = {
                    "SupplierName": supplier_name,
                    "SupplierID": supplier_id,
                    "InvoiceNumber": invoice_number,
                    "RawMaterial": raw_material,
                    "RawMaterialCode": raw_material_code,
                    "QuantityPurchased": 0.0,
                    "QuantityConsumed": 0.0,
                    "QuantityAvailable": 0.0,
                }

            source_map[key]["QuantityPurchased"] += quantity_purchased

        sources = []

        for key, source in source_map.items():
            consumed = consumption.get(key, 0.0)
            available = source["QuantityPurchased"] - consumed

            source["QuantityConsumed"] = round(consumed, 4)
            source["QuantityAvailable"] = round(max(available, 0.0), 4)

            if source["QuantityAvailable"] > 0:
                sources.append(source)

        sources.sort(
            key=lambda item: (
                item["SupplierName"].casefold(),
                item["InvoiceNumber"].casefold(),
                item["RawMaterialCode"].casefold(),
            )
        )

        return {
            "next_batch_id": generate_next_batch_id(workbook),
            "sources": sources,
        }

    finally:
        workbook.close()


# ============================================================
# MAIN
# ============================================================

try:
    print(json.dumps(read_sources(), indent=4))
except Exception as error:
    fail(str(error))

