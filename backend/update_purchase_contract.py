import json
import os
import sys
import warnings
from datetime import date, datetime
from decimal import Decimal, InvalidOperation

from openpyxl import load_workbook

warnings.filterwarnings("ignore")


# =========================================================
# CONFIG
# =========================================================
# =============================================
# CONFIG PATH
# =============================================

config_path = os.path.join(
    BASE_DIR,
    "Config",
    "database_config.json"
)
SHEET_NAME = "pcon"
START_ROW = 5


# =========================================================
# SAFE ERROR HANDLER
# =========================================================

def fail(message):
    sys.stderr.write(f"ERROR: {message}\n")
    sys.exit(1)


# =========================================================
# DATABASE CONFIG
# =========================================================

def get_database_path():
    if not os.path.exists(CONFIG_PATH):
        fail("database_config.json not found.")

    try:
        with open(CONFIG_PATH, "r", encoding="utf-8") as file:
            config = json.load(file)
    except json.JSONDecodeError:
        fail("Invalid JSON inside database_config.json.")
    except Exception as error:
        fail(str(error))

    database_path = (
        config.get("database_path")
        or config.get("excel_path")
        or config.get("path")
        or ""
    ).strip()

    if database_path == "":
        fail("Database path is empty.")

    if not os.path.exists(database_path):
        fail("Database file not found.")

    return database_path


# =========================================================
# SAFE VALUE FUNCTIONS
# =========================================================

def safe_string(value):
    if value is None:
        return ""

    return str(value).strip()


def parse_decimal(value, field_name):
    try:
        text = (
            safe_string(value)
            .replace(",", "")
            .replace("%", "")
        )

        if text == "":
            fail(f"{field_name} is required.")

        return float(Decimal(text))

    except (InvalidOperation, ValueError):
        fail(f"{field_name} must be numeric.")


def number_or_zero(value):
    try:
        if value is None:
            return 0

        text = (
            str(value)
            .strip()
            .replace(",", "")
            .replace("%", "")
        )

        if text == "":
            return 0

        return float(Decimal(text))

    except Exception:
        return 0


def parse_date(value):
    if value is None:
        return None

    if isinstance(value, datetime):
        return value.date()

    if isinstance(value, date):
        return value

    text = str(value).strip()

    if text == "":
        return None

    for date_format in (
        "%d-%m-%Y",
        "%d/%m/%Y",
        "%Y-%m-%d",
        "%m/%d/%Y",
    ):
        try:
            return datetime.strptime(
                text,
                date_format
            ).date()
        except ValueError:
            continue

    return None


def is_yes(value):
    text = safe_string(value).upper()

    return text in {
        "YES",
        "Y",
        "TRUE",
        "1",
        "BREACH",
        "BREACHED"
    }


# =========================================================
# STATUS CHECK
# =========================================================

def calculate_status(
    start_date,
    end_date,
    remaining_qty,
    breach_detected,
    renewal_reference
):
    today = date.today()

    if safe_string(renewal_reference) != "":
        return "Expired"

    if remaining_qty <= 0:
        if end_date is not None and today > end_date:
            return "Expired"

        return "Completed"

    if start_date is not None and today < start_date:
        return "Upcoming"

    if end_date is not None and today > end_date and remaining_qty > 0:
        return "Violated"

    if breach_detected and remaining_qty > 0:
        return "Violated"

    return "Active"


# =========================================================
# ARGUMENT VALIDATION
# =========================================================

EXPECTED_ARGUMENTS = 9

if len(sys.argv) != EXPECTED_ARGUMENTS:
    fail("Invalid number of arguments received from C#.")

try:
    row_number = int(sys.argv[1])
except Exception:
    fail("Invalid row number.")

vendor_name = safe_string(sys.argv[2])
vendor_id = safe_string(sys.argv[3])

base_price = parse_decimal(
    sys.argv[4],
    "Base Price"
)

agreed_qty = parse_decimal(
    sys.argv[5],
    "Agreed Quantity"
)

penalty_percent = parse_decimal(
    sys.argv[6],
    "Penalty Percentage"
)

remedy_days = parse_decimal(
    sys.argv[7],
    "Remedy Days"
)

breach_responsibility = safe_string(
    sys.argv[8]
)


# =========================================================
# BUSINESS VALIDATION
# =========================================================

if vendor_name == "":
    fail("Vendor Name cannot be empty.")

if vendor_id == "":
    fail("Vendor ID cannot be empty.")

if base_price <= 0:
    fail("Base Price must be greater than zero.")

if agreed_qty <= 0:
    fail("Agreed Quantity must be greater than zero.")

if penalty_percent < 0:
    fail("Penalty Percentage cannot be negative.")

if penalty_percent > 100:
    fail("Penalty Percentage cannot be greater than 100.")

if penalty_percent > 1:
    penalty_percent = penalty_percent / 100

if remedy_days < 0:
    fail("Remedy Days cannot be negative.")

ALLOWED_RESPONSIBILITIES = {
    "Vendor",
    "Company",
    "Shared",
    "Pending"
}

if breach_responsibility == "" or breach_responsibility == "None":
    breach_responsibility = "Pending"

if breach_responsibility not in ALLOWED_RESPONSIBILITIES:
    fail("Invalid breach responsibility value.")


# =========================================================
# LOAD WORKBOOKS
# =========================================================

DATABASE_PATH = get_database_path()

try:
    values_workbook = load_workbook(
        DATABASE_PATH,
        data_only=True,
        read_only=True
    )
except PermissionError:
    fail("Close Excel workbook before continuing.")
except Exception as error:
    fail(f"Failed to open workbook for validation.\n{str(error)}")

try:
    workbook = load_workbook(
        DATABASE_PATH,
        data_only=False
    )
except PermissionError:
    values_workbook.close()
    fail("Close Excel workbook before continuing.")
except Exception as error:
    values_workbook.close()
    fail(f"Failed to open workbook.\n{str(error)}")


# =========================================================
# VALIDATE SHEETS
# =========================================================

try:
    if SHEET_NAME not in workbook.sheetnames:
        fail(f"Sheet '{SHEET_NAME}' does not exist.")

    if SHEET_NAME not in values_workbook.sheetnames:
        fail(f"Sheet '{SHEET_NAME}' does not exist.")

    sheet = workbook[SHEET_NAME]
    values_sheet = values_workbook[SHEET_NAME]

    if row_number < START_ROW:
        fail("Invalid contract row.")

    if row_number > sheet.max_row:
        fail("Contract row does not exist.")

    if safe_string(values_sheet[f"B{row_number}"].value) == "":
        fail("No contract exists in selected row.")

    renewal_reference = safe_string(
        values_sheet[f"Y{row_number}"].value
    )

    start_date = parse_date(
        values_sheet[f"F{row_number}"].value
    )

    end_date = parse_date(
        values_sheet[f"G{row_number}"].value
    )

    remaining_qty = number_or_zero(
        values_sheet[f"O{row_number}"].value
    )

    breach_detected = is_yes(
        values_sheet[f"Q{row_number}"].value
    )

    dynamic_status = calculate_status(
        start_date,
        end_date,
        remaining_qty,
        breach_detected,
        renewal_reference
    )

    if dynamic_status == "Expired":
        fail("Expired historical contracts cannot be edited.")

    if not breach_detected and breach_responsibility != "Pending":
        fail(
            "Breach responsibility can only be assigned when breach is detected."
        )

    # =====================================================
    # UPDATE EDITABLE CELLS ONLY
    # =====================================================

    sheet[f"B{row_number}"] = vendor_name
    sheet[f"C{row_number}"] = vendor_id

    sheet[f"L{row_number}"] = base_price
    sheet[f"M{row_number}"] = agreed_qty

    sheet[f"S{row_number}"] = penalty_percent
    sheet[f"S{row_number}"].number_format = "0%"

    sheet[f"V{row_number}"] = remedy_days

    if breach_detected:
        sheet[f"R{row_number}"] = breach_responsibility
    elif safe_string(sheet[f"R{row_number}"].value) == "":
        sheet[f"R{row_number}"] = "Pending"

    try:
        workbook.calculation.fullCalcOnLoad = True
        workbook.calculation.forceFullCalc = True
    except Exception:
        pass

    workbook.save(DATABASE_PATH)

    print("Contract updated successfully.")

except PermissionError:
    fail("Cannot save workbook. Close Excel file first.")

except Exception as error:
    fail(str(error))

finally:
    values_workbook.close()
    workbook.close()