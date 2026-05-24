import json
import os
import sys
from copy import copy
from datetime import datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.formula.translate import Translator


# ============================================================
# CONFIG
# ============================================================

CONFIG_PATH = (
    Path.home()
    / "Documents"
    / "RubexOps"
    / "database_config.json"
)

SHEET_NAME = "scon"
START_ROW = 5

DEFAULT_ITEM_NAME = "Natural Rubber Field Coagulum"


# ============================================================
# COLUMN MAP
# ============================================================

COL_SERIAL_NO = "A"
COL_CUSTOMER_NAME = "B"
COL_CUSTOMER_ID = "C"
COL_ITEM_NAME = "D"
COL_ITEM_CODE = "E"
COL_START_DATE = "F"
COL_END_DATE = "G"
COL_DAYS_REMAINING = "H"
COL_SALES_SO_FAR = "I"
COL_BASE_PRICE = "L"
COL_AGREED_QTY = "M"
COL_SOLD_QTY = "N"
COL_REMAINING_QTY = "O"
COL_COMPLETION_PERCENT = "P"
COL_BREACH = "Q"
COL_BREACH_RESPONSIBILITY = "R"
COL_PENALTY_PERCENT = "S"
COL_REVISED_RATE = "T"
COL_PENALTY_VALUE = "U"
COL_REMEDY_DAYS = "V"
COL_REMEDY_DEADLINE = "W"
COL_REMEDY_STATUS = "X"
COL_RENEWAL_REFERENCE = "Y"


# ============================================================
# SAFE ERROR HANDLER
# ============================================================

def fail(message):
    print(f"ERROR: {message}", file=sys.stderr)
    sys.exit(1)


# ============================================================
# CONFIG HELPERS
# ============================================================

def get_database_path():
    if not CONFIG_PATH.exists():
        fail(f"Database config file not found: {CONFIG_PATH}")

    try:
        with open(CONFIG_PATH, "r", encoding="utf-8") as file:
            config = json.load(file)
    except json.JSONDecodeError:
        fail("Invalid JSON inside database_config.json.")

    database_path = (
        config.get("database_path")
        or config.get("excel_path")
        or config.get("path")
        or ""
    ).strip()

    if not database_path:
        fail("Database path was not found in database_config.json.")

    if not os.path.exists(database_path):
        fail(f"Excel database file not found: {database_path}")

    return database_path


# ============================================================
# VALIDATION HELPERS
# ============================================================

def safe_string(value):
    if value is None:
        return ""

    return str(value).strip()


def parse_date(value, field_name):
    text = safe_string(value)

    if not text:
        fail(f"{field_name} is required.")

    try:
        return datetime.strptime(text, "%d-%m-%Y")
    except ValueError:
        fail(f"{field_name} must be in dd-MM-yyyy format.")


def parse_decimal(value, field_name, required=True):
    text = safe_string(value).replace(",", "").replace("%", "")

    if not text:
        if required:
            fail(f"{field_name} is required.")

        return 0.0

    try:
        return float(Decimal(text))
    except (InvalidOperation, ValueError):
        fail(f"{field_name} must be numeric.")


# ============================================================
# ROW HELPERS
# ============================================================

def get_next_contract_row(sheet):
    row = START_ROW

    while True:
        customer_name = sheet[f"{COL_CUSTOMER_NAME}{row}"].value
        customer_id = sheet[f"{COL_CUSTOMER_ID}{row}"].value

        if safe_string(customer_name) == "" and safe_string(customer_id) == "":
            return row

        row += 1


def copy_row_style_and_formulas(sheet, source_row, target_row):
    for column in range(1, sheet.max_column + 1):
        source_cell = sheet.cell(row=source_row, column=column)
        target_cell = sheet.cell(row=target_row, column=column)

        if source_cell.has_style:
            target_cell.font = copy(source_cell.font)
            target_cell.fill = copy(source_cell.fill)
            target_cell.border = copy(source_cell.border)
            target_cell.alignment = copy(source_cell.alignment)
            target_cell.number_format = source_cell.number_format
            target_cell.protection = copy(source_cell.protection)

        if isinstance(source_cell.value, str) and source_cell.value.startswith("="):
            target_cell.value = Translator(
                source_cell.value,
                origin=source_cell.coordinate
            ).translate_formula(target_cell.coordinate)

    if source_row in sheet.row_dimensions:
        sheet.row_dimensions[target_row].height = (
            sheet.row_dimensions[source_row].height
        )


def ensure_sales_formulas(sheet, row):
    sheet[f"{COL_DAYS_REMAINING}{row}"] = (
        f'=IF({COL_END_DATE}{row}="","",{COL_END_DATE}{row}-TODAY())'
    )

    sheet[f"{COL_SALES_SO_FAR}{row}"] = (
        f'=IF({COL_CUSTOMER_ID}{row}="","",'
        f'COUNTIFS(sales!$C:$C,${COL_CUSTOMER_ID}{row},'
        f'sales!$E:$E,${COL_ITEM_CODE}{row},'
        f'sales!$H:$H,">="&${COL_START_DATE}{row},'
        f'sales!$H:$H,"<="&${COL_END_DATE}{row}))'
    )

    sheet[f"{COL_SOLD_QTY}{row}"] = (
        f'=IF({COL_CUSTOMER_ID}{row}="","",'
        f'SUMIFS(sales!$I:$I,'
        f'sales!$C:$C,${COL_CUSTOMER_ID}{row},'
        f'sales!$E:$E,${COL_ITEM_CODE}{row},'
        f'sales!$H:$H,">="&${COL_START_DATE}{row},'
        f'sales!$H:$H,"<="&${COL_END_DATE}{row}))'
    )

    sheet[f"{COL_REMAINING_QTY}{row}"] = (
        f'=IF({COL_AGREED_QTY}{row}="","",{COL_AGREED_QTY}{row}-{COL_SOLD_QTY}{row})'
    )

    sheet[f"{COL_COMPLETION_PERCENT}{row}"] = (
        f'=IFERROR(({COL_SOLD_QTY}{row}/{COL_AGREED_QTY}{row})*100,0)'
    )

    sheet[f"{COL_BREACH}{row}"] = (
        f'=IF(OR({COL_END_DATE}{row}="",{COL_REMAINING_QTY}{row}=""),"",'
        f'IF(AND({COL_END_DATE}{row}<TODAY(),{COL_REMAINING_QTY}{row}>0),"YES","NO"))'
    )

    sheet[f"{COL_REMEDY_DEADLINE}{row}"] = (
        f'=IF(AND({COL_BREACH}{row}="YES",{COL_REMEDY_DAYS}{row}<>""),'
        f'{COL_END_DATE}{row}+{COL_REMEDY_DAYS}{row},"")'
    )

    sheet[f"{COL_REMEDY_STATUS}{row}"] = (
        f'=IF({COL_BREACH}{row}<>"YES","",'
        f'IF({COL_REMAINING_QTY}{row}<=0,"CLEARED",'
        f'IF(TODAY()<={COL_REMEDY_DEADLINE}{row},"IN REMEDY","OVERDUE")))'
    )


# ============================================================
# CREATE SALES CONTRACT
# ============================================================

def create_sales_contract(args):
    if len(args) != 9:
        fail(
            "Expected 9 arguments: customer name, customer id, item code, "
            "start date, end date, base price, agreed quantity, "
            "penalty percentage, remedy days."
        )

    (
        customer_name,
        customer_id,
        item_code,
        start_date_text,
        end_date_text,
        base_price_text,
        agreed_qty_text,
        penalty_percent_text,
        remedy_days_text,
    ) = args

    customer_name = safe_string(customer_name)
    customer_id = safe_string(customer_id)
    item_code = safe_string(item_code) or "NRFC"

    if not customer_name:
        fail("Customer Name is required.")

    if not customer_id:
        fail("Customer ID is required.")

    start_date = parse_date(start_date_text, "Start Date")
    end_date = parse_date(end_date_text, "End Date")

    if end_date <= start_date:
        fail("End Date must be after Start Date.")

    base_price = parse_decimal(base_price_text, "Base Price")
    agreed_qty = parse_decimal(agreed_qty_text, "Agreed Quantity")
    penalty_percent = parse_decimal(
        penalty_percent_text,
        "Penalty Percentage",
        required=False
    )
    remedy_days = parse_decimal(
        remedy_days_text,
        "Remedy Days",
        required=False
    )

    if base_price <= 0:
        fail("Base Price must be greater than zero.")

    if agreed_qty <= 0:
        fail("Agreed Quantity must be greater than zero.")

    if penalty_percent < 0 or penalty_percent > 100:
        fail("Penalty Percentage must be between 0 and 100.")

    if penalty_percent > 1:
        penalty_percent = penalty_percent / 100

    if remedy_days < 0:
        fail("Remedy Days cannot be negative.")

    database_path = get_database_path()

    try:
        workbook = load_workbook(database_path, data_only=False)
    except PermissionError:
        fail("Close Excel workbook before continuing.")
    except Exception as error:
        fail(f"Failed to open workbook.\n{error}")

    try:
        if SHEET_NAME not in workbook.sheetnames:
            fail(f"Sheet not found: {SHEET_NAME}")

        sheet = workbook[SHEET_NAME]

        target_row = get_next_contract_row(sheet)
        template_row = target_row - 1 if target_row > START_ROW else START_ROW

        copy_row_style_and_formulas(sheet, template_row, target_row)

        sheet[f"{COL_SERIAL_NO}{target_row}"] = target_row - START_ROW + 1
        sheet[f"{COL_CUSTOMER_NAME}{target_row}"] = customer_name
        sheet[f"{COL_CUSTOMER_ID}{target_row}"] = customer_id
        sheet[f"{COL_ITEM_NAME}{target_row}"] = DEFAULT_ITEM_NAME
        sheet[f"{COL_ITEM_CODE}{target_row}"] = item_code
        sheet[f"{COL_START_DATE}{target_row}"] = start_date
        sheet[f"{COL_END_DATE}{target_row}"] = end_date
        sheet[f"{COL_START_DATE}{target_row}"].number_format = "dd-mm-yyyy"
        sheet[f"{COL_END_DATE}{target_row}"].number_format = "dd-mm-yyyy"
        sheet[f"{COL_BASE_PRICE}{target_row}"] = base_price
        sheet[f"{COL_AGREED_QTY}{target_row}"] = agreed_qty
        sheet[f"{COL_BREACH_RESPONSIBILITY}{target_row}"] = "Pending"
        sheet[f"{COL_PENALTY_PERCENT}{target_row}"] = penalty_percent
        sheet[f"{COL_PENALTY_PERCENT}{target_row}"].number_format = "0%"
        sheet[f"{COL_REMEDY_DAYS}{target_row}"] = remedy_days
        sheet[f"{COL_RENEWAL_REFERENCE}{target_row}"] = None

        ensure_sales_formulas(sheet, target_row)

        workbook.calculation.fullCalcOnLoad = True
        workbook.calculation.forceFullCalc = True
        workbook.save(database_path)

        return f"Sales contract created successfully at row {target_row}."

    except PermissionError:
        fail("Cannot save workbook. Close Excel file first.")
    finally:
        workbook.close()


# ============================================================
# MAIN
# ============================================================

try:
    print(create_sales_contract(sys.argv[1:]))
except Exception as error:
    fail(str(error))

