import json
import os
import sys
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import load_workbook


# ============================================================
# CONFIG
# ============================================================

CONFIG_PATH = (
    Path.home()
    / "Documents"
    / "RubexOps"
    / "database_config.json"
)

SHEET_NAME = "scon"
START_ROW = 5


# ============================================================
# SAFE ERROR HANDLER
# ============================================================

def fail(message):
    print(f"ERROR: {message}", file=sys.stderr)
    sys.exit(1)


# ============================================================
# HELPERS
# ============================================================

def safe_string(value):
    if value is None:
        return ""

    return str(value).strip()


def get_database_path():
    if not CONFIG_PATH.exists():
        fail("database_config.json not found.")

    try:
        with open(CONFIG_PATH, "r", encoding="utf-8") as file:
            config = json.load(file)
    except json.JSONDecodeError:
        fail("Invalid JSON inside database_config.json.")

    database_path = (
        config.get("database_path")
        or config.get("excel_path")
        or config.get("path")
        or ""
    ).strip()

    if not database_path:
        fail("Database path is empty.")

    if not os.path.exists(database_path):
        fail("Database file not found.")

    return database_path


def parse_decimal(value, field_name):
    text = safe_string(value).replace(",", "").replace("%", "")

    if not text:
        fail(f"{field_name} is required.")

    try:
        return float(Decimal(text))
    except (InvalidOperation, ValueError):
        fail(f"{field_name} must be numeric.")


def number_or_zero(value):
    try:
        text = safe_string(value).replace(",", "").replace("%", "")

        if not text:
            return 0

        return float(Decimal(text))
    except Exception:
        return 0


def parse_date(value):
    if value is None:
        return None

    if isinstance(value, datetime):
        return value.date()

    if isinstance(value, date):
        return value

    text = str(value).strip()

    if not text:
        return None

    for date_format in ("%d-%m-%Y", "%d/%m/%Y", "%Y-%m-%d", "%m/%d/%Y"):
        try:
            return datetime.strptime(text, date_format).date()
        except ValueError:
            continue

    return None


def is_yes(value):
    return safe_string(value).upper() in {"YES", "Y", "TRUE", "1", "BREACH", "BREACHED"}


def calculate_status(start_date, end_date, remaining_qty, breach_detected, renewal_reference):
    today = date.today()

    if safe_string(renewal_reference) != "":
        return "Expired"

    if remaining_qty <= 0:
        if end_date is not None and today > end_date:
            return "Expired"

        return "Completed"

    if start_date is not None and today < start_date:
        return "Upcoming"

    if breach_detected:
        return "Violated"

    if end_date is not None and today > end_date and remaining_qty > 0:
        return "Violated"

    return "Active"


# ============================================================
# ARGUMENTS
# ============================================================

if len(sys.argv) != 9:
    fail("Invalid number of arguments received from C#.")

try:
    row_number = int(sys.argv[1])
except ValueError:
    fail("Invalid row number.")

customer_name = safe_string(sys.argv[2])
customer_id = safe_string(sys.argv[3])
base_price = parse_decimal(sys.argv[4], "Base Price")
agreed_qty = parse_decimal(sys.argv[5], "Agreed Quantity")
penalty_percent = parse_decimal(sys.argv[6], "Penalty / Discount Percentage")
remedy_days = parse_decimal(sys.argv[7], "Remedy Days")
breach_responsibility = safe_string(sys.argv[8])

if not customer_name:
    fail("Customer Name cannot be empty.")

if not customer_id:
    fail("Customer ID cannot be empty.")

if base_price <= 0:
    fail("Base Price must be greater than zero.")

if agreed_qty <= 0:
    fail("Agreed Quantity must be greater than zero.")

if penalty_percent < 0 or penalty_percent > 100:
    fail("Penalty / Discount Percentage must be between 0 and 100.")

if penalty_percent > 1:
    penalty_percent = penalty_percent / 100

if remedy_days < 0:
    fail("Remedy Days cannot be negative.")

allowed_responsibilities = {"Pending", "Company", "Customer", "Shared"}

if breach_responsibility == "":
    breach_responsibility = "Pending"

if breach_responsibility not in allowed_responsibilities:
    fail("Invalid breach responsibility value.")


# ============================================================
# UPDATE WORKBOOK
# ============================================================

database_path = get_database_path()

try:
    values_workbook = load_workbook(database_path, data_only=True, read_only=True)
    workbook = load_workbook(database_path, data_only=False)
except PermissionError:
    fail("Close Excel workbook before continuing.")
except Exception as error:
    fail(f"Failed to open workbook.\n{error}")

try:
    if SHEET_NAME not in workbook.sheetnames or SHEET_NAME not in values_workbook.sheetnames:
        fail(f"Sheet '{SHEET_NAME}' does not exist.")

    sheet = workbook[SHEET_NAME]
    values_sheet = values_workbook[SHEET_NAME]

    if row_number < START_ROW or row_number > sheet.max_row:
        fail("Invalid contract row.")

    if safe_string(values_sheet[f"B{row_number}"].value) == "" and safe_string(values_sheet[f"C{row_number}"].value) == "":
        fail("No contract exists in selected row.")

    start_date = parse_date(values_sheet[f"F{row_number}"].value)
    end_date = parse_date(values_sheet[f"G{row_number}"].value)
    remaining_qty = number_or_zero(values_sheet[f"O{row_number}"].value)
    breach_detected = is_yes(values_sheet[f"Q{row_number}"].value)
    renewal_reference = safe_string(values_sheet[f"Y{row_number}"].value)

    dynamic_status = calculate_status(
        start_date,
        end_date,
        remaining_qty,
        breach_detected,
        renewal_reference
    )

    if dynamic_status == "Expired":
        fail("Expired historical sales contracts cannot be edited.")

    if not breach_detected and breach_responsibility != "Pending":
        fail("Breach responsibility can only be assigned when breach is detected.")

    sheet[f"B{row_number}"] = customer_name
    sheet[f"C{row_number}"] = customer_id
    sheet[f"L{row_number}"] = base_price
    sheet[f"M{row_number}"] = agreed_qty
    sheet[f"S{row_number}"] = penalty_percent
    sheet[f"S{row_number}"].number_format = "0%"
    sheet[f"V{row_number}"] = remedy_days

    if breach_detected:
        sheet[f"R{row_number}"] = breach_responsibility
    elif safe_string(sheet[f"R{row_number}"].value) == "":
        sheet[f"R{row_number}"] = "Pending"

    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    workbook.save(database_path)

    print("Sales contract updated successfully.")

except PermissionError:
    fail("Cannot save workbook. Close Excel file first.")
except Exception as error:
    fail(str(error))
finally:
    values_workbook.close()
    workbook.close()

