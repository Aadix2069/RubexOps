import json
import os
import sys
from copy import copy
from datetime import datetime, date
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.formula.translate import Translator


# ============================================================
# CONFIG
# ============================================================

CONFIG_PATH = (
    Path.home()
    / "Documents"
    / "RubexOps"
    / "database_config.json"
)

SHEET_NAME = "scon"
START_ROW = 5


# ============================================================
# COLUMN MAP
# ============================================================

COL_CUSTOMER_NAME = "B"
COL_CUSTOMER_ID = "C"
COL_ITEM_NAME = "D"
COL_ITEM_CODE = "E"
COL_START_DATE = "F"
COL_END_DATE = "G"
COL_BASE_PRICE = "L"
COL_AGREED_QTY = "M"
COL_BREACH_RESPONSIBILITY = "R"
COL_PENALTY_PERCENT = "S"
COL_REMEDY_DAYS = "V"
COL_RENEWAL_REFERENCE = "Y"


# ============================================================
# ERROR HANDLER
# ============================================================

def fail(message):
    print(f"ERROR: {message}", file=sys.stderr)
    sys.exit(1)


# ============================================================
# HELPERS
# ============================================================

def safe_string(value):
    if value is None:
        return ""

    return str(value).strip()


def get_database_path():
    if not CONFIG_PATH.exists():
        fail("database_config.json not found.")

    try:
        with open(CONFIG_PATH, "r", encoding="utf-8") as file:
            config = json.load(file)
    except json.JSONDecodeError:
        fail("Invalid JSON inside database_config.json.")

    database_path = (
        config.get("database_path")
        or config.get("excel_path")
        or config.get("path")
        or ""
    ).strip()

    if not database_path:
        fail("Database path is empty.")

    if not os.path.exists(database_path):
        fail("Database file not found.")

    return database_path


def parse_date_text(value, field_name):
    text = safe_string(value)

    if not text:
        fail(f"{field_name} is required.")

    try:
        return datetime.strptime(text, "%d-%m-%Y").date()
    except ValueError:
        fail(f"{field_name} must be in dd-MM-yyyy format.")


def parse_existing_date(value):
    if isinstance(value, datetime):
        return value.date()

    if isinstance(value, date):
        return value

    text = safe_string(value)

    for date_format in ("%d-%m-%Y", "%d/%m/%Y", "%Y-%m-%d", "%m/%d/%Y"):
        try:
            return datetime.strptime(text, date_format).date()
        except ValueError:
            continue

    return None


def parse_decimal(value, field_name):
    text = safe_string(value).replace(",", "").replace("%", "")

    if not text:
        fail(f"{field_name} is required.")

    try:
        return float(Decimal(text))
    except (InvalidOperation, ValueError):
        fail(f"{field_name} must be numeric.")


def get_next_row(sheet):
    row = START_ROW

    while True:
        customer_name = sheet[f"{COL_CUSTOMER_NAME}{row}"].value
        customer_id = sheet[f"{COL_CUSTOMER_ID}{row}"].value

        if safe_string(customer_name) == "" and safe_string(customer_id) == "":
            return row

        row += 1


def copy_row_style_and_formulas(sheet, source_row, target_row):
    for column in range(1, sheet.max_column + 1):
        source_cell = sheet.cell(row=source_row, column=column)
        target_cell = sheet.cell(row=target_row, column=column)

        if source_cell.has_style:
            target_cell.font = copy(source_cell.font)
            target_cell.fill = copy(source_cell.fill)
            target_cell.border = copy(source_cell.border)
            target_cell.alignment = copy(source_cell.alignment)
            target_cell.number_format = source_cell.number_format
            target_cell.protection = copy(source_cell.protection)

        if isinstance(source_cell.value, str) and source_cell.value.startswith("="):
            target_cell.value = Translator(
                source_cell.value,
                origin=source_cell.coordinate
            ).translate_formula(target_cell.coordinate)
        else:
            target_cell.value = None

    if source_row in sheet.row_dimensions:
        sheet.row_dimensions[target_row].height = sheet.row_dimensions[source_row].height


# ============================================================
# RENEW SALES CONTRACT
# ============================================================

def renew_sales_contract(args):
    if len(args) != 7:
        fail(
            "Expected 7 arguments: old row, new start date, new end date, "
            "base price, agreed quantity, penalty percentage, remedy days."
        )

    try:
        old_row = int(args[0])
    except ValueError:
        fail("Invalid old contract row.")

    new_start_date = parse_date_text(args[1], "New Start Date")
    new_end_date = parse_date_text(args[2], "New End Date")

    if new_end_date <= new_start_date:
        fail("New End Date must be after New Start Date.")

    base_price = parse_decimal(args[3], "Base Price")
    agreed_qty = parse_decimal(args[4], "Agreed Quantity")
    penalty_percent = parse_decimal(args[5], "Penalty / Discount Percentage")
    remedy_days = parse_decimal(args[6], "Remedy Days")

    if base_price <= 0:
        fail("Base Price must be greater than zero.")

    if agreed_qty <= 0:
        fail("Agreed Quantity must be greater than zero.")

    if penalty_percent < 0 or penalty_percent > 100:
        fail("Penalty / Discount Percentage must be between 0 and 100.")

    if penalty_percent > 1:
        penalty_percent = penalty_percent / 100

    if remedy_days < 0:
        fail("Remedy Days cannot be negative.")

    database_path = get_database_path()

    try:
        workbook = load_workbook(database_path, data_only=False)
    except PermissionError:
        fail("Close Excel workbook before continuing.")
    except Exception as error:
        fail(f"Failed to open workbook.\n{error}")

    try:
        if SHEET_NAME not in workbook.sheetnames:
            fail(f"Sheet '{SHEET_NAME}' does not exist.")

        sheet = workbook[SHEET_NAME]

        if old_row < START_ROW or old_row > sheet.max_row:
            fail("Invalid old contract row.")

        if safe_string(sheet[f"{COL_CUSTOMER_ID}{old_row}"].value) == "":
            fail("No sales contract exists in selected row.")

        renewal_reference = safe_string(sheet[f"{COL_RENEWAL_REFERENCE}{old_row}"].value)

        if renewal_reference:
            fail(f"This sales contract has already been renewed into row {renewal_reference}.")

        old_end_date = parse_existing_date(sheet[f"{COL_END_DATE}{old_row}"].value)

        if old_end_date is not None and new_start_date <= old_end_date:
            fail("New Start Date must be after the old contract End Date to prevent sales overlap.")

        new_row = get_next_row(sheet)

        copy_row_style_and_formulas(sheet, old_row, new_row)

        sheet[f"{COL_CUSTOMER_NAME}{new_row}"] = sheet[f"{COL_CUSTOMER_NAME}{old_row}"].value
        sheet[f"{COL_CUSTOMER_ID}{new_row}"] = sheet[f"{COL_CUSTOMER_ID}{old_row}"].value
        sheet[f"{COL_ITEM_NAME}{new_row}"] = sheet[f"{COL_ITEM_NAME}{old_row}"].value
        sheet[f"{COL_ITEM_CODE}{new_row}"] = sheet[f"{COL_ITEM_CODE}{old_row}"].value
        sheet[f"{COL_START_DATE}{new_row}"] = new_start_date
        sheet[f"{COL_END_DATE}{new_row}"] = new_end_date
        sheet[f"{COL_START_DATE}{new_row}"].number_format = "dd-mm-yyyy"
        sheet[f"{COL_END_DATE}{new_row}"].number_format = "dd-mm-yyyy"
        sheet[f"{COL_BASE_PRICE}{new_row}"] = base_price
        sheet[f"{COL_AGREED_QTY}{new_row}"] = agreed_qty
        sheet[f"{COL_BREACH_RESPONSIBILITY}{new_row}"] = "Pending"
        sheet[f"{COL_PENALTY_PERCENT}{new_row}"] = penalty_percent
        sheet[f"{COL_PENALTY_PERCENT}{new_row}"].number_format = "0%"
        sheet[f"{COL_REMEDY_DAYS}{new_row}"] = remedy_days
        sheet[f"{COL_RENEWAL_REFERENCE}{old_row}"] = new_row
        sheet[f"{COL_RENEWAL_REFERENCE}{new_row}"] = None

        workbook.calculation.fullCalcOnLoad = True
        workbook.calculation.forceFullCalc = True
        workbook.save(database_path)

        return (
            f"Sales contract renewed successfully. "
            f"Old row {old_row} is now historical. New contract created at row {new_row}."
        )

    except PermissionError:
        fail("Cannot save workbook. Close Excel file first.")
    finally:
        workbook.close()


try:
    print(renew_sales_contract(sys.argv[1:]))
except Exception as error:
    fail(str(error))

