import json
import os
import sys
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import load_workbook


# ============================================================
# CONFIG
# ============================================================

CONFIG_PATH = (
    Path.home()
    / "Documents"
    / "RubexOps"
    / "database_config.json"
)

SCON_SHEET_NAME = "scon"
SALES_SHEET_NAME = "sales"
START_ROW = 5


# ============================================================
# COLUMN MAP
# ============================================================

COL_CUSTOMER_NAME = "B"
COL_CUSTOMER_ID = "C"
COL_ITEM_NAME = "D"
COL_ITEM_CODE = "E"
COL_START_DATE = "F"
COL_END_DATE = "G"
COL_DAYS_REMAINING = "H"
COL_SALES_SO_FAR = "I"
COL_BASE_PRICE = "L"
COL_AGREED_QTY = "M"
COL_SOLD_QTY = "N"
COL_REMAINING_QTY = "O"
COL_COMPLETION_PERCENT = "P"
COL_BREACH = "Q"
COL_BREACH_RESPONSIBILITY = "R"
COL_PENALTY_PERCENT = "S"
COL_REVISED_RATE = "T"
COL_PENALTY_VALUE = "U"
COL_REMEDY_DAYS = "V"
COL_REMEDY_DEADLINE = "W"
COL_REMEDY_STATUS = "X"
COL_RENEWAL_REFERENCE = "Y"

SALES_COL_CUSTOMER_ID = "C"
SALES_COL_ITEM_CODE = "E"
SALES_COL_DISPATCH_DATE = "H"
SALES_COL_WEIGHT = "I"


# ============================================================
# SAFE ERROR HANDLER
# ============================================================

def fail(message):
    print(f"ERROR: {message}", file=sys.stderr)
    sys.exit(1)


# ============================================================
# CONFIG HELPERS
# ============================================================

def get_database_path():
    if not CONFIG_PATH.exists():
        fail("database_config.json not found.")

    try:
        with open(CONFIG_PATH, "r", encoding="utf-8") as file:
            config = json.load(file)
    except json.JSONDecodeError:
        fail("Invalid JSON inside database_config.json.")

    database_path = (
        config.get("database_path")
        or config.get("excel_path")
        or config.get("path")
        or ""
    ).strip()

    if not database_path:
        fail("Database path is empty.")

    if not os.path.exists(database_path):
        fail("Database file not found.")

    return database_path


# ============================================================
# SAFE VALUE HELPERS
# ============================================================

def safe_string(value):
    if value is None:
        return ""

    return str(value).strip()


def number_or_none(value):
    if value is None or isinstance(value, bool):
        return None

    try:
        if isinstance(value, str):
            text = value.strip().replace(",", "").replace("%", "")

            if not text:
                return None

            return float(Decimal(text))

        return float(value)
    except (ValueError, TypeError, InvalidOperation):
        return None


def safe_number(value):
    number = number_or_none(value)

    if number is None:
        return 0

    return round(number, 4)


def safe_percent(value):
    number = number_or_none(value)

    if number is None:
        return 0

    if 0 < abs(number) <= 1:
        number = number * 100

    return round(number, 4)


def parse_date(value):
    if value is None:
        return None

    if isinstance(value, datetime):
        return value.date()

    if isinstance(value, date):
        return value

    text = str(value).strip()

    if not text:
        return None

    for date_format in ("%d-%m-%Y", "%d/%m/%Y", "%Y-%m-%d", "%m/%d/%Y"):
        try:
            return datetime.strptime(text, date_format).date()
        except ValueError:
            continue

    return None


def safe_date(value):
    parsed_date = parse_date(value)

    if parsed_date is None:
        return safe_string(value)

    return parsed_date.strftime("%d-%m-%Y")


def is_yes(value):
    text = safe_string(value).upper()

    return text in {"YES", "Y", "TRUE", "1", "BREACH", "BREACHED"}


# ============================================================
# SALES FALLBACK CALCULATION
# ============================================================

def build_sales_records(workbook):
    if SALES_SHEET_NAME not in workbook.sheetnames:
        return []

    sheet = workbook[SALES_SHEET_NAME]
    records = []

    for row in range(START_ROW, sheet.max_row + 1):
        customer_id = safe_string(sheet[f"{SALES_COL_CUSTOMER_ID}{row}"].value)
        item_code = safe_string(sheet[f"{SALES_COL_ITEM_CODE}{row}"].value)
        dispatch_date = parse_date(sheet[f"{SALES_COL_DISPATCH_DATE}{row}"].value)

        if not customer_id or dispatch_date is None:
            continue

        sold_qty = number_or_none(sheet[f"{SALES_COL_WEIGHT}{row}"].value) or 0

        records.append(
            {
                "customer_id": customer_id.casefold(),
                "item_code": item_code.casefold(),
                "dispatch_date": dispatch_date,
                "sold_qty": sold_qty,
            }
        )

    return records


def calculate_sales_totals(sales_records, customer_id, item_code, start_date, end_date):
    if not customer_id or start_date is None or end_date is None:
        return {
            "sales_so_far": 0,
            "sold_qty": 0,
        }

    customer_key = customer_id.casefold()
    item_key = item_code.casefold()

    matches = []

    for record in sales_records:
        if record["customer_id"] != customer_key:
            continue

        if item_key and record["item_code"] and record["item_code"] != item_key:
            continue

        if start_date <= record["dispatch_date"] <= end_date:
            matches.append(record)

    return {
        "sales_so_far": len(matches),
        "sold_qty": sum(record["sold_qty"] for record in matches),
    }


# ============================================================
# STATUS ENGINE
# ============================================================

def calculate_status(start_date, end_date, remaining_qty, breach_detected, renewal_reference):
    today = date.today()

    if safe_string(renewal_reference) != "":
        return "Expired", "#F59E0B"

    if remaining_qty <= 0:
        if end_date is not None and today > end_date:
            return "Expired", "#F59E0B"

        return "Completed", "#8B5CF6"

    if start_date is not None and today < start_date:
        return "Upcoming", "#3B82F6"

    if breach_detected:
        return "Violated", "#EF4444"

    if end_date is not None and today > end_date and remaining_qty > 0:
        return "Violated", "#EF4444"

    return "Active", "#10B981"


# ============================================================
# MAIN READ
# ============================================================

database_path = get_database_path()

try:
    workbook = load_workbook(database_path, data_only=True)
except PermissionError:
    fail("Close Excel workbook before continuing.")
except Exception as error:
    fail(f"Failed to load workbook.\n{error}")

try:
    if SCON_SHEET_NAME not in workbook.sheetnames:
        fail(f"Sheet not found: {SCON_SHEET_NAME}")

    sheet = workbook[SCON_SHEET_NAME]
    sales_records = build_sales_records(workbook)
    contracts = []

    for row in range(START_ROW, sheet.max_row + 1):
        customer_name = safe_string(sheet[f"{COL_CUSTOMER_NAME}{row}"].value)
        customer_id = safe_string(sheet[f"{COL_CUSTOMER_ID}{row}"].value)

        if customer_name == "" and customer_id == "":
            continue

        item_code = safe_string(sheet[f"{COL_ITEM_CODE}{row}"].value)
        start_date = parse_date(sheet[f"{COL_START_DATE}{row}"].value)
        end_date = parse_date(sheet[f"{COL_END_DATE}{row}"].value)

        agreed_qty = safe_number(sheet[f"{COL_AGREED_QTY}{row}"].value)

        excel_sold_qty = number_or_none(sheet[f"{COL_SOLD_QTY}{row}"].value)
        excel_remaining_qty = number_or_none(sheet[f"{COL_REMAINING_QTY}{row}"].value)
        excel_completion_percent = number_or_none(sheet[f"{COL_COMPLETION_PERCENT}{row}"].value)
        excel_sales_so_far = number_or_none(sheet[f"{COL_SALES_SO_FAR}{row}"].value)

        sales_totals = calculate_sales_totals(
            sales_records,
            customer_id,
            item_code,
            start_date,
            end_date
        )

        sold_qty = excel_sold_qty if excel_sold_qty is not None else 0

        if sales_totals["sold_qty"] > 0 and sold_qty == 0:
            sold_qty = sales_totals["sold_qty"]

        calculated_remaining_qty = max(agreed_qty - sold_qty, 0)

        if excel_remaining_qty is None:
            remaining_qty = calculated_remaining_qty
        elif sales_totals["sold_qty"] > 0 and abs(excel_remaining_qty - calculated_remaining_qty) > 0.01:
            remaining_qty = calculated_remaining_qty
        else:
            remaining_qty = excel_remaining_qty

        calculated_completion_percent = (
            min((sold_qty / agreed_qty) * 100, 100)
            if agreed_qty > 0
            else 0
        )

        if excel_completion_percent is None:
            completion_percent = calculated_completion_percent
        else:
            completion_percent = safe_percent(excel_completion_percent)

            if sales_totals["sold_qty"] > 0 and completion_percent == 0 and calculated_completion_percent > 0:
                completion_percent = calculated_completion_percent

        sales_so_far = int(excel_sales_so_far) if excel_sales_so_far is not None else sales_totals["sales_so_far"]

        if sales_so_far == 0 and sales_totals["sales_so_far"] > 0:
            sales_so_far = sales_totals["sales_so_far"]

        breach_detected = is_yes(sheet[f"{COL_BREACH}{row}"].value)
        renewal_reference = safe_string(sheet[f"{COL_RENEWAL_REFERENCE}{row}"].value)

        dynamic_status, status_color = calculate_status(
            start_date,
            end_date,
            remaining_qty,
            breach_detected,
            renewal_reference
        )

        is_expired = dynamic_status == "Expired"

        contracts.append(
            {
                "row": row,
                "customer_name": customer_name,
                "customer_id": customer_id,
                "item_name": safe_string(sheet[f"{COL_ITEM_NAME}{row}"].value),
                "item_code": item_code,
                "start_date": safe_date(sheet[f"{COL_START_DATE}{row}"].value),
                "end_date": safe_date(sheet[f"{COL_END_DATE}{row}"].value),
                "days_remaining": safe_string(sheet[f"{COL_DAYS_REMAINING}{row}"].value),
                "sales_so_far": str(sales_so_far),
                "base_price": safe_number(sheet[f"{COL_BASE_PRICE}{row}"].value),
                "agreed_qty": agreed_qty,
                "sold_qty": round(sold_qty, 4),
                "remaining_qty": round(remaining_qty, 4),
                "completion_percent": round(completion_percent, 4),
                "breach": "YES" if breach_detected else "NO",
                "breach_responsibility": safe_string(sheet[f"{COL_BREACH_RESPONSIBILITY}{row}"].value),
                "penalty_percent": safe_percent(sheet[f"{COL_PENALTY_PERCENT}{row}"].value),
                "revised_rate": safe_number(sheet[f"{COL_REVISED_RATE}{row}"].value),
                "penalty_value": safe_number(sheet[f"{COL_PENALTY_VALUE}{row}"].value),
                "remedy_days": safe_number(sheet[f"{COL_REMEDY_DAYS}{row}"].value),
                "remedy_deadline": safe_date(sheet[f"{COL_REMEDY_DEADLINE}{row}"].value),
                "remedy_status": safe_string(sheet[f"{COL_REMEDY_STATUS}{row}"].value),
                "renewal_reference": renewal_reference,
                "status": dynamic_status,
                "dynamic_status": dynamic_status,
                "status_color": status_color,
                "is_expired": is_expired,
                "can_edit": not is_expired,
                "can_renew": (
                    renewal_reference == ""
                    and dynamic_status in ["Completed", "Violated", "Expired"]
                ),
                "is_breach_detected": breach_detected,
                "show_responsibility": breach_detected,
                "show_breach_panel": breach_detected,
            }
        )

    contracts.sort(
        key=lambda item: (
            item["is_expired"],
            {
                "Active": 0,
                "Upcoming": 1,
                "Violated": 2,
                "Completed": 3,
                "Expired": 4,
            }.get(item["dynamic_status"], 5),
            item["customer_name"].lower(),
        )
    )

    print(json.dumps(contracts, indent=4))

except Exception as error:
    fail(f"Failed to read sales contracts.\n{error}")
finally:
    workbook.close()

