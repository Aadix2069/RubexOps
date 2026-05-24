import json
import os
import sys
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import load_workbook


# ============================================================
# CONFIG
# ============================================================

CONFIG_PATH = (
    Path.home()
    / "Documents"
    / "RubexOps"
    / "database_config.json"
)

SHEET_NAME = "production"
START_ROW = 5


# ============================================================
# ERROR HANDLER
# ============================================================

def fail(message):
    print(f"ERROR: {message}", file=sys.stderr)
    sys.exit(1)


# ============================================================
# SAFE HELPERS
# ============================================================

def safe_string(value):
    if value is None:
        return ""

    return str(value).strip()


def number_or_zero(value):
    if value is None:
        return 0.0

    try:
        if isinstance(value, str):
            text = value.strip().replace(",", "").replace("%", "")

            if not text:
                return 0.0

            return float(Decimal(text))

        return float(value)
    except (InvalidOperation, ValueError, TypeError):
        return 0.0


def parse_date(value):
    if value is None:
        return None

    if isinstance(value, datetime):
        return value.date()

    if isinstance(value, date):
        return value

    text = str(value).strip()

    if not text:
        return None

    for date_format in ("%d-%m-%Y", "%d/%m/%Y", "%Y-%m-%d", "%m/%d/%Y"):
        try:
            return datetime.strptime(text, date_format).date()
        except ValueError:
            continue

    return None


def safe_date(value):
    parsed_date = parse_date(value)

    if parsed_date is None:
        return safe_string(value)

    return parsed_date.strftime("%d-%m-%Y")


def get_database_path():
    if not CONFIG_PATH.exists():
        fail("database_config.json not found.")

    try:
        with open(CONFIG_PATH, "r", encoding="utf-8") as file:
            config = json.load(file)
    except json.JSONDecodeError:
        fail("Invalid JSON inside database_config.json.")

    database_path = (
        config.get("database_path")
        or config.get("excel_path")
        or config.get("path")
        or ""
    ).strip()

    if not database_path:
        fail("Database path is empty.")

    if not os.path.exists(database_path):
        fail("Database file not found.")

    return database_path


# ============================================================
# READ RECORDS
# ============================================================

def read_records():
    database_path = get_database_path()

    try:
        workbook = load_workbook(database_path, data_only=True, read_only=True)
    except PermissionError:
        fail("Close Excel workbook before continuing.")
    except Exception as error:
        fail(f"Failed to open workbook.\n{error}")

    try:
        if SHEET_NAME not in workbook.sheetnames:
            fail(f"Sheet not found: {SHEET_NAME}")

        sheet = workbook[SHEET_NAME]
        records = []

        for row in range(START_ROW, sheet.max_row + 1):
            batch_id = safe_string(sheet[f"B{row}"].value)
            invoice_number = safe_string(sheet[f"F{row}"].value)

            if not batch_id and not invoice_number:
                continue

            input_weight = number_or_zero(sheet[f"L{row}"].value)
            output_weight = number_or_zero(sheet[f"M{row}"].value)
            production_loss = number_or_zero(sheet[f"O{row}"].value)
            drc_variance = number_or_zero(sheet[f"R{row}"].value)

            loss_percent = (
                (production_loss / input_weight) * 100
                if input_weight > 0
                else 0
            )

            records.append(
                {
                    "row": row,
                    "batch_id": batch_id,
                    "production_date": safe_date(sheet[f"C{row}"].value),
                    "supplier_name": safe_string(sheet[f"D{row}"].value),
                    "supplier_id": safe_string(sheet[f"E{row}"].value),
                    "invoice_number": invoice_number,
                    "raw_material": safe_string(sheet[f"G{row}"].value),
                    "raw_material_code": safe_string(sheet[f"H{row}"].value),
                    "finished_product": safe_string(sheet[f"I{row}"].value),
                    "finished_product_code": safe_string(sheet[f"J{row}"].value),
                    "quantity_available": number_or_zero(sheet[f"K{row}"].value),
                    "input_weight": input_weight,
                    "output_weight": output_weight,
                    "remaining_quantity": number_or_zero(sheet[f"N{row}"].value),
                    "production_loss": production_loss,
                    "initial_drc": number_or_zero(sheet[f"P{row}"].value),
                    "actual_drc": number_or_zero(sheet[f"Q{row}"].value),
                    "drc_variance": drc_variance,
                    "loss_percent": round(loss_percent, 4),
                    "is_high_loss": loss_percent >= 10,
                    "is_high_drc_variance": abs(drc_variance) >= 5,
                }
            )

        today_text = date.today().strftime("%d-%m-%Y")
        today_records = [
            record for record in records
            if record["production_date"] == today_text
        ]

        total_output = sum(record["output_weight"] for record in records)
        total_loss = sum(record["production_loss"] for record in records)
        average_drc_variance = (
            sum(record["drc_variance"] for record in records) / len(records)
            if records
            else 0
        )

        highest_loss_record = max(
            records,
            key=lambda item: item["production_loss"],
            default=None
        )

        return {
            "records": records,
            "summary": {
                "total_production_today": round(
                    sum(record["output_weight"] for record in today_records),
                    4
                ),
                "total_output": round(total_output, 4),
                "total_production_loss": round(total_loss, 4),
                "average_drc_variance": round(average_drc_variance, 4),
                "active_batches": len(records),
                "highest_loss_batch": (
                    highest_loss_record["batch_id"]
                    if highest_loss_record
                    else ""
                ),
            },
        }

    finally:
        workbook.close()


# ============================================================
# MAIN
# ============================================================

try:
    print(json.dumps(read_records(), indent=4))
except Exception as error:
    fail(str(error))

