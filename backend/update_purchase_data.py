import json
import os
import sys
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook


# ============================================================
# CONFIG
# ============================================================

CONFIG_PATH = (
    Path.home()
    / "Documents"
    / "RubexOps"
    / "database_config.json"
)

PCON_SHEET_NAME = "pcon"
PURCHASE_SHEET_NAME = "purchase"
START_ROW = 5
DATE_FORMAT = "%d-%m-%Y"


# ============================================================
# COLUMN MAP
# ============================================================

COL_VENDOR_NAME = "B"
COL_VENDOR_ID = "C"
COL_ITEM_NAME = "D"
COL_ITEM_CODE = "E"
COL_INVOICE_NUMBER = "F"
COL_PURCHASE_ORDER = "G"
COL_DELIVERY = "H"
COL_INVOICE_WEIGHT = "I"
COL_BEFORE_UNLOADING = "J"
COL_CARRIER_WEIGHT = "K"
COL_NO_OF_BAGS = "M"
COL_CALCULATED_DRC = "O"
COL_GST_PERCENT = "S"
COL_TDS_194Q_PERCENT = "T"
COL_UNLOADING_CHARGE = "V"

DEFAULT_ITEM_NAME = "Natural Rubber Field Coagulum"


# ============================================================
# DATABASE PATH
# ============================================================

def get_database_path():
    if not CONFIG_PATH.exists():
        raise FileNotFoundError(
            f"Database config file not found: {CONFIG_PATH}"
        )

    with open(CONFIG_PATH, "r", encoding="utf-8") as file:
        config = json.load(file)

    database_path = (
        config.get("database_path")
        or config.get("excel_path")
        or config.get("path")
    )

    if not database_path:
        raise ValueError(
            "Database path was not found in database_config.json."
        )

    if not os.path.exists(database_path):
        raise FileNotFoundError(
            f"Excel database file not found: {database_path}"
        )

    return database_path


# ============================================================
# VALUE HELPERS
# ============================================================

def clean_number(value):
    return str(value).strip().replace("%", "").replace(",", "")


def parse_required_decimal(value, field_name, allow_zero):
    if value is None or str(value).strip() == "":
        raise ValueError(f"{field_name} is required.")

    try:
        number = float(clean_number(value))
    except ValueError as error:
        raise ValueError(
            f"{field_name} must be a valid number."
        ) from error

    if allow_zero:
        if number < 0:
            raise ValueError(f"{field_name} cannot be negative.")
    elif number <= 0:
        raise ValueError(f"{field_name} must be greater than zero.")

    return number


def parse_optional_decimal(value, field_name, allow_zero):
    if value is None or str(value).strip() == "":
        return None

    try:
        number = float(clean_number(value))
    except ValueError as error:
        raise ValueError(
            f"{field_name} must be a valid number."
        ) from error

    if allow_zero:
        if number < 0:
            raise ValueError(f"{field_name} cannot be negative.")
    elif number <= 0:
        raise ValueError(
            f"{field_name} must be greater than zero when entered."
        )

    return number


def parse_required_percent(value, field_name, allow_zero):
    number = parse_required_decimal(
        value,
        field_name,
        allow_zero
    )

    if number > 100:
        raise ValueError(f"{field_name} cannot be greater than 100.")

    return number


def parse_optional_integer(value, field_name):
    if value is None or str(value).strip() == "":
        return None

    try:
        number = int(str(value).strip())
    except ValueError as error:
        raise ValueError(
            f"{field_name} must be a whole number."
        ) from error

    if number < 0:
        raise ValueError(f"{field_name} cannot be negative.")

    return number


def parse_date(value, field_name):
    if value is None or str(value).strip() == "":
        raise ValueError(f"{field_name} is required.")

    try:
        return datetime.strptime(
            str(value).strip(),
            DATE_FORMAT
        )
    except ValueError as error:
        raise ValueError(
            f"{field_name} must be in dd-MM-yyyy format."
        ) from error


# ============================================================
# BUSINESS VALIDATION
# ============================================================

def validate_vendor(workbook, vendor_name, vendor_id):
    if PCON_SHEET_NAME not in workbook.sheetnames:
        raise ValueError(f"Sheet not found: {PCON_SHEET_NAME}")

    sheet = workbook[PCON_SHEET_NAME]
    vendor_id_text = str(vendor_id).strip().lower()
    vendor_name_text = str(vendor_name).strip().lower()

    for row in range(START_ROW, sheet.max_row + 1):
        sheet_vendor_name = sheet[f"B{row}"].value
        sheet_vendor_id = sheet[f"C{row}"].value

        if sheet_vendor_name is None or sheet_vendor_id is None:
            continue

        if str(sheet_vendor_id).strip().lower() != vendor_id_text:
            continue

        if str(sheet_vendor_name).strip().lower() != vendor_name_text:
            raise ValueError(
                "Vendor Name and Vendor ID do not match the pcon sheet."
            )

        return

    raise ValueError(
        "Vendor ID was not found in the pcon sheet."
    )


def validate_purchase_row(sheet, row_number):
    if row_number < START_ROW:
        raise ValueError(
            "Invalid purchase row selected."
        )

    if row_number > sheet.max_row:
        raise ValueError(
            "Selected purchase row does not exist."
        )

    vendor_id = sheet[f"{COL_VENDOR_ID}{row_number}"].value
    invoice_number = sheet[f"{COL_INVOICE_NUMBER}{row_number}"].value

    if vendor_id is None and invoice_number is None:
        raise ValueError(
            "Selected purchase row is empty."
        )


def validate_invoice_number(sheet, row_number, invoice_number):
    invoice_text = str(invoice_number).strip().lower()

    for row in range(START_ROW, sheet.max_row + 1):
        if row == row_number:
            continue

        existing_invoice = sheet[f"{COL_INVOICE_NUMBER}{row}"].value

        if existing_invoice is None:
            continue

        if str(existing_invoice).strip().lower() == invoice_text:
            raise ValueError(
                "Invoice Number already exists in the purchase sheet."
            )


# ============================================================
# UPDATE PURCHASE DATA
# ============================================================

def update_purchase_data(args):
    if len(args) != 15:
        raise ValueError(
            "Expected 15 arguments: row number, vendor name, vendor id, "
            "item code, invoice number, purchase order date, delivery date, "
            "invoice weight, before unloading, carrier weight, no. of bags, "
            "calculated DRC, GST, TDS 194Q, unloading charge."
        )

    (
        row_number,
        vendor_name,
        vendor_id,
        item_code,
        invoice_number,
        purchase_order_date,
        delivery_date,
        invoice_weight,
        before_unloading,
        carrier_weight,
        no_of_bags,
        calculated_drc,
        gst_percent,
        tds_194q_percent,
        unloading_charge,
    ) = args

    try:
        row_number = int(row_number)
    except ValueError as error:
        raise ValueError(
            "Invalid purchase row number."
        ) from error

    if not str(vendor_name).strip():
        raise ValueError("Vendor Name is required.")

    if not str(vendor_id).strip():
        raise ValueError("Vendor ID is required.")

    if not str(invoice_number).strip():
        raise ValueError("Invoice Number is required.")

    if not str(item_code).strip():
        item_code = "NRFC"

    purchase_order_value = parse_date(
        purchase_order_date,
        "Purchase Order date"
    )

    delivery_value = parse_date(
        delivery_date,
        "Delivery date"
    )

    if delivery_value < purchase_order_value:
        raise ValueError(
            "Delivery date must be on or after Purchase Order date."
        )

    invoice_weight_value = parse_optional_decimal(
        invoice_weight,
        "Invoice Weight",
        False
    )

    before_unloading_value = parse_required_decimal(
        before_unloading,
        "Before Unloading",
        False
    )

    carrier_weight_value = parse_required_decimal(
        carrier_weight,
        "Carrier Weight",
        False
    )

    if carrier_weight_value > before_unloading_value:
        raise ValueError(
            "Carrier Weight cannot be greater than Before Unloading weight."
        )

    no_of_bags_value = parse_optional_integer(
        no_of_bags,
        "No. of Bags"
    )

    calculated_drc_value = parse_required_percent(
        calculated_drc,
        "Calculated DRC",
        False
    )

    gst_percent_value = parse_required_percent(
        gst_percent,
        "GST",
        True
    )

    tds_194q_percent_value = parse_required_percent(
        tds_194q_percent,
        "TDS 194Q",
        True
    )

    unloading_charge_value = parse_required_decimal(
        unloading_charge,
        "Unloading Charge",
        True
    )

    database_path = get_database_path()

    workbook = load_workbook(
        database_path,
        data_only=False
    )

    if PURCHASE_SHEET_NAME not in workbook.sheetnames:
        raise ValueError(
            f"Sheet not found: {PURCHASE_SHEET_NAME}"
        )

    validate_vendor(
        workbook,
        vendor_name,
        vendor_id
    )

    sheet = workbook[PURCHASE_SHEET_NAME]

    validate_purchase_row(
        sheet,
        row_number
    )

    validate_invoice_number(
        sheet,
        row_number,
        invoice_number
    )

    sheet[f"{COL_VENDOR_NAME}{row_number}"] = vendor_name
    sheet[f"{COL_VENDOR_ID}{row_number}"] = vendor_id
    sheet[f"{COL_ITEM_NAME}{row_number}"] = DEFAULT_ITEM_NAME
    sheet[f"{COL_ITEM_CODE}{row_number}"] = item_code
    sheet[f"{COL_INVOICE_NUMBER}{row_number}"] = invoice_number
    sheet[f"{COL_PURCHASE_ORDER}{row_number}"] = purchase_order_date
    sheet[f"{COL_DELIVERY}{row_number}"] = delivery_date
    sheet[f"{COL_INVOICE_WEIGHT}{row_number}"] = invoice_weight_value
    sheet[f"{COL_BEFORE_UNLOADING}{row_number}"] = before_unloading_value
    sheet[f"{COL_CARRIER_WEIGHT}{row_number}"] = carrier_weight_value
    sheet[f"{COL_NO_OF_BAGS}{row_number}"] = no_of_bags_value
    sheet[f"{COL_CALCULATED_DRC}{row_number}"] = calculated_drc_value
    sheet[f"{COL_GST_PERCENT}{row_number}"] = gst_percent_value
    sheet[f"{COL_TDS_194Q_PERCENT}{row_number}"] = tds_194q_percent_value
    sheet[f"{COL_UNLOADING_CHARGE}{row_number}"] = unloading_charge_value

    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True

    workbook.save(database_path)

    return f"Purchase data updated successfully at row {row_number}."


# ============================================================
# MAIN
# ============================================================

try:
    print(
        update_purchase_data(
            sys.argv[1:]
        )
    )

except Exception as error:
    print(f"ERROR: {error}")
    sys.exit(1)
