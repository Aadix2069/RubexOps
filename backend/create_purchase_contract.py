import sys
import os
import json
import warnings
from openpyxl import load_workbook
from datetime import datetime

warnings.filterwarnings("ignore")



# =========================================================
# SAFE ERROR FUNCTION
# =========================================================

def fail(message):

    # IMPORTANT:
    # SEND ERRORS TO STDERR
    # SO JSON / OUTPUT DOES NOT BREAK

    sys.stderr.write(f"ERROR: {message}\n")

    sys.exit(1)



# =========================================================
# LOAD DATABASE CONFIG
# =========================================================

try:

    # =============================================
    # GET BASE DIRECTORY
    # =============================================

    if getattr(sys, 'frozen', False):

        BASE_DIR = os.path.dirname(
            sys.executable
        )

    else:

        BASE_DIR = os.path.dirname(
            os.path.abspath(__file__)
        )



    # =============================================
    # CONFIG PATH
    # =============================================
# =============================================
# CONFIG PATH
# =============================================

    config_path = os.path.join(
    BASE_DIR,
    "Config",
    "database_config.json"
)



    # =============================================
    # CHECK CONFIG EXISTS
    # =============================================

    if not os.path.exists(config_path):

        fail("database_config.json not found.")



    # =============================================
    # READ CONFIG
    # =============================================

    with open(config_path, "r") as file:

        config = json.load(file)



except json.JSONDecodeError:

    fail("Invalid JSON inside database_config.json.")

except Exception as ex:

    fail(str(ex))



# =========================================================
# DATABASE PATH
# =========================================================

DATABASE_PATH = config.get(
    "database_path",
    ""
).strip()



SHEET_NAME = "pcon"



# =========================================================
# VALIDATE DATABASE PATH
# =========================================================

if DATABASE_PATH == "":

    fail("Database path is empty.")



if not os.path.exists(DATABASE_PATH):

    fail("Database file not found.")



# =========================================================
# COLUMN MAPPING
# =========================================================

COLUMN_MAP = {
    "vendor_name": "B",
    "vendor_id": "C",
    "item_code": "E",
    "start_date": "F",
    "end_date": "G",
    "base_price": "L",
    "agreed_quantity": "M",
    "penalty_rate": "S",
    "remedy_days": "V"
}



# =========================================================
# VALIDATE ARGUMENT COUNT
# =========================================================

EXPECTED_ARGUMENTS = 10



if len(sys.argv) != EXPECTED_ARGUMENTS:

    fail("Missing required arguments.")



# =========================================================
# SAFE CONVERSION FUNCTIONS
# =========================================================

def safe_string(value):

    if value is None:

        return ""



    return str(value).strip()



def safe_float(value, field_name):

    try:

        number = float(value)

        return number

    except:

        fail(f"{field_name} must be numeric.")



# =========================================================
# GET ARGUMENTS
# =========================================================

vendor_name = safe_string(
    sys.argv[1]
)



vendor_id = safe_string(
    sys.argv[2]
)



item_code = safe_string(
    sys.argv[3]
)



start_date = safe_string(
    sys.argv[4]
)



end_date = safe_string(
    sys.argv[5]
)



base_price = safe_float(
    sys.argv[6],
    "Base Price"
)



agreed_quantity = safe_float(
    sys.argv[7],
    "Agreed Quantity"
)



penalty_rate = safe_float(
    sys.argv[8],
    "Penalty Rate"
)



remedy_days = safe_float(
    sys.argv[9],
    "Remedy Days"
)



# =========================================================
# EMPTY VALIDATION
# =========================================================

if vendor_name == "":

    fail("Vendor Name cannot be empty.")



if vendor_id == "":

    fail("Vendor ID cannot be empty.")



if item_code == "":

    fail("Item Code cannot be empty.")



# =========================================================
# NEGATIVE VALIDATION
# =========================================================

if base_price <= 0:

    fail("Base Price must be greater than 0.")



if agreed_quantity <= 0:

    fail("Agreed Quantity must be greater than 0.")



if penalty_rate < 0:

    fail("Penalty Rate cannot be negative.")



if remedy_days < 0:

    fail("Remedy Days cannot be negative.")



# =========================================================
# DATE VALIDATION
# =========================================================

try:

    start_date_object = datetime.strptime(
        start_date,
        "%d-%m-%Y"
    )



    end_date_object = datetime.strptime(
        end_date,
        "%d-%m-%Y"
    )



except:

    fail("Invalid date format.")



if end_date_object <= start_date_object:

    fail(
        "End date must be after start date."
    )



# =========================================================
# LOAD WORKBOOK
# =========================================================

try:

    workbook = load_workbook(
        DATABASE_PATH
    )



except PermissionError:

    fail(
        "Close Excel workbook before continuing."
    )

except Exception as ex:

    fail(
        f"Failed to load workbook.\n{str(ex)}"
    )



# =========================================================
# VALIDATE SHEET
# =========================================================

if SHEET_NAME not in workbook.sheetnames:

    fail(
        f"Sheet '{SHEET_NAME}' does not exist."
    )



sheet = workbook[SHEET_NAME]



# =========================================================
# DUPLICATE VENDOR ID VALIDATION
# =========================================================

for row in range(5, sheet.max_row + 1):

    existing_vendor_id = sheet[f"C{row}"].value

    termination_status = sheet[f"Y{row}"].value



    if existing_vendor_id is None:

        continue



    existing_vendor_id = str(
        existing_vendor_id
    ).strip().lower()



    current_vendor_id = vendor_id.strip().lower()



    # =============================================
    # SKIP TERMINATED CONTRACTS
    # =============================================

    is_terminated = False



    if termination_status is not None:

        is_terminated = (
            str(termination_status)
            .strip()
            .upper()
            == "TERMINATED"
        )



    # =============================================
    # DUPLICATE CHECK
    # =============================================

    if (
        existing_vendor_id == current_vendor_id
        and not is_terminated
    ):

        fail(
            "Vendor ID already exists.\n"
            "Please use a different Vendor ID."
        )



# =========================================================
# FIND NEXT EMPTY ROW
# =========================================================

next_row = 5



while sheet[f"B{next_row}"].value not in [None, ""]:

    next_row += 1



# =========================================================
# WRITE DATA
# =========================================================

try:

    # =============================================
    # BASIC DETAILS
    # =============================================

    sheet[
        f"{COLUMN_MAP['vendor_name']}{next_row}"
    ] = vendor_name



    sheet[
        f"{COLUMN_MAP['vendor_id']}{next_row}"
    ] = vendor_id



    sheet[
        f"{COLUMN_MAP['item_code']}{next_row}"
    ] = item_code



    # =============================================
    # DATES
    # =============================================

    sheet[
        f"{COLUMN_MAP['start_date']}{next_row}"
    ] = start_date_object



    sheet[
        f"{COLUMN_MAP['end_date']}{next_row}"
    ] = end_date_object



    sheet[
        f"{COLUMN_MAP['start_date']}{next_row}"
    ].number_format = "DD-MM-YYYY"



    sheet[
        f"{COLUMN_MAP['end_date']}{next_row}"
    ].number_format = "DD-MM-YYYY"



    # =============================================
    # CONTRACT VALUES
    # =============================================

    sheet[
        f"{COLUMN_MAP['base_price']}{next_row}"
    ] = base_price



    sheet[
        f"{COLUMN_MAP['agreed_quantity']}{next_row}"
    ] = agreed_quantity



    # =============================================
    # PENALTY RATE
    # =============================================

    if penalty_rate > 1:

        penalty_rate = penalty_rate / 100



    sheet[
        f"{COLUMN_MAP['penalty_rate']}{next_row}"
    ] = penalty_rate



    sheet[
        f"{COLUMN_MAP['penalty_rate']}{next_row}"
    ].number_format = "0%"



    # =============================================
    # REMEDY DAYS
    # =============================================

    sheet[
        f"{COLUMN_MAP['remedy_days']}{next_row}"
    ] = remedy_days



    # =============================================
    # DEFAULT VALUES
    # =============================================

    sheet[f"R{next_row}"] = "Pending"



    # =============================================
    # CLEAR TERMINATION STATUS
    # =============================================

    sheet[f"Y{next_row}"] = ""



except Exception as ex:

    fail(
        f"Failed to write contract data.\n{str(ex)}"
    )



# =========================================================
# SAVE WORKBOOK
# =========================================================

try:

    workbook.save(
        DATABASE_PATH
    )



except PermissionError:

    fail(
        "Cannot save workbook. Close Excel file first."
    )

except Exception as ex:

    fail(
        f"Failed to save workbook.\n{str(ex)}"
    )



# =========================================================
# SUCCESS
# =========================================================

print(
    "Purchase Contract Created Successfully"
)