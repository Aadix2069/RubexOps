import json
import os
import sys
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel


# ============================================================
# CONFIG
# ============================================================

CONFIG_PATH = (
    Path.home()
    / "Documents"
    / "RubexOps"
    / "database_config.json"
)

PCON_SHEET_NAME = "pcon"
PURCHASE_SHEET_NAME = "purchase"
START_ROW = 5
DATE_FORMAT = "%d-%m-%Y"


# ============================================================
# COLUMN MAP
# ============================================================

COL_SERIAL_NO = "A"
COL_VENDOR_NAME = "B"
COL_VENDOR_ID = "C"
COL_ITEM_NAME = "D"
COL_ITEM_CODE = "E"
COL_INVOICE_NUMBER = "F"
COL_PURCHASE_ORDER = "G"
COL_DELIVERY = "H"
COL_INVOICE_WEIGHT = "I"
COL_BEFORE_UNLOADING = "J"
COL_CARRIER_WEIGHT = "K"
COL_RECEIVED_WEIGHT = "L"
COL_NO_OF_BAGS = "M"
COL_NET_WEIGHT = "N"
COL_CALCULATED_DRC = "O"
COL_DRC_WEIGHT = "P"
COL_BASE_RATE = "Q"
COL_ADJUSTED_RATE = "R"
COL_GST_PERCENT = "S"
COL_TDS_194Q_PERCENT = "T"
COL_TAXABLE_AMOUNT = "U"
COL_UNLOADING_CHARGE = "V"
COL_GST_AMOUNT = "W"
COL_GROSS_AMOUNT = "X"
COL_TDS_AMOUNT = "Y"
COL_NET_PAYABLE = "Z"


# ============================================================
# DATABASE PATH
# ============================================================

def get_database_path():
    if not CONFIG_PATH.exists():
        raise FileNotFoundError(
            f"Database config file not found: {CONFIG_PATH}"
        )

    with open(CONFIG_PATH, "r", encoding="utf-8") as file:
        config = json.load(file)

    database_path = (
        config.get("database_path")
        or config.get("excel_path")
        or config.get("path")
    )

    if not database_path:
        raise ValueError(
            "Database path was not found in database_config.json."
        )

    if not os.path.exists(database_path):
        raise FileNotFoundError(
            f"Excel database file not found: {database_path}"
        )

    return database_path


# ============================================================
# VALUE HELPERS
# ============================================================

def cell_value(sheet, column, row):
    return sheet[f"{column}{row}"].value


def text_value(value):
    if value is None:
        return ""

    return str(value).strip()


def number_value(value):
    if value is None or value == "":
        return None

    if isinstance(value, (int, float)):
        return round(float(value), 4)

    text = str(value).strip().replace("%", "").replace(",", "")

    if text == "":
        return None

    try:
        return round(float(text), 4)
    except ValueError:
        return None


def int_value(value):
    number = number_value(value)

    if number is None:
        return None

    return int(number)


def date_value(value):
    parsed = parse_date_value(value)

    if parsed is None:
        return text_value(value)

    return parsed.strftime(DATE_FORMAT)


def parse_date_value(value):
    if value is None or value == "":
        return None

    if isinstance(value, datetime):
        return value

    if isinstance(value, (int, float)):
        try:
            return from_excel(value)
        except Exception:
            return None

    text = str(value).strip()

    for date_format in (
        "%d-%m-%Y",
        "%d/%m/%Y",
        "%Y-%m-%d",
        "%m/%d/%Y",
    ):
        try:
            return datetime.strptime(text, date_format)
        except ValueError:
            continue

    return None


def fallback(value, calculated_value):
    if value is not None:
        return value

    return calculated_value


def round_money(value):
    if value is None:
        return None

    return round(float(value), 2)


# ============================================================
# CONTRACT RATE FALLBACK
# ============================================================

def find_contract_base_rate(workbook, vendor_id, delivery_date):
    contract_info = find_contract_info(
        workbook,
        vendor_id,
        delivery_date
    )

    if contract_info is None:
        return None

    return contract_info.get("BaseRate")


def contract_rank(contract, delivery_date):
    start_date = contract["StartDate"]
    end_date = contract["EndDate"]

    if delivery_date and start_date and end_date:
        if start_date.date() <= delivery_date.date() <= end_date.date():
            return (0, end_date)

    today = datetime.today()

    if start_date and end_date and start_date.date() <= today.date() <= end_date.date():
        return (1, end_date)

    if start_date and today.date() < start_date.date():
        return (2, start_date)

    if end_date:
        return (3, end_date)

    return (4, datetime.min)


def choose_contract(existing, candidate, delivery_date):
    if existing is None:
        return candidate

    existing_rank = contract_rank(existing, delivery_date)
    candidate_rank = contract_rank(candidate, delivery_date)

    if candidate_rank[0] < existing_rank[0]:
        return candidate

    if candidate_rank[0] == existing_rank[0]:
        if candidate_rank[1] and existing_rank[1]:
            if candidate_rank[1] > existing_rank[1]:
                return candidate

    return existing


def find_contract_info(workbook, vendor_id, delivery_date):
    if PCON_SHEET_NAME not in workbook.sheetnames:
        return None

    if not vendor_id:
        return None

    sheet = workbook[PCON_SHEET_NAME]
    vendor_id_text = str(vendor_id).strip().lower()
    matched_contract = None

    for row in range(START_ROW, sheet.max_row + 1):
        contract_vendor_id = cell_value(sheet, "C", row)

        if contract_vendor_id is None:
            continue

        if str(contract_vendor_id).strip().lower() != vendor_id_text:
            continue

        start_date = parse_date_value(
            cell_value(sheet, "F", row)
        )

        end_date = parse_date_value(
            cell_value(sheet, "G", row)
        )

        contract = {
            "VendorName": text_value(cell_value(sheet, "B", row)),
            "VendorID": text_value(contract_vendor_id),
            "ItemName": text_value(cell_value(sheet, "D", row)),
            "ItemCode": text_value(cell_value(sheet, "E", row)),
            "BaseRate": number_value(cell_value(sheet, "L", row)),
            "StartDate": start_date,
            "EndDate": end_date,
        }

        matched_contract = choose_contract(
            matched_contract,
            contract,
            delivery_date
        )

    return matched_contract


# ============================================================
# FORMULA FALLBACKS
# ============================================================

def apply_calculated_fallbacks(workbook, row_data):
    before_unloading = row_data["BeforeUnloading"]
    carrier_weight = row_data["CarrierWeight"]
    no_of_bags = row_data["NoOfBags"] or 0
    calculated_drc = row_data["CalculatedDrc"]
    gst_percent = row_data["GstPercent"]
    tds_percent = row_data["Tds194QPercent"]
    unloading_charge = row_data["UnloadingCharge"] or 0

    if before_unloading is not None and carrier_weight is not None:
        row_data["ReceivedWeight"] = fallback(
            row_data["ReceivedWeight"],
            round(before_unloading - carrier_weight, 4)
        )

    if row_data["ReceivedWeight"] is not None:
        row_data["NetWeight"] = fallback(
            row_data["NetWeight"],
            round(row_data["ReceivedWeight"] - no_of_bags, 4)
        )

    if row_data["NetWeight"] is not None and calculated_drc is not None:
        row_data["DrcWeight"] = fallback(
            row_data["DrcWeight"],
            round(row_data["NetWeight"] * calculated_drc / 100, 4)
        )

    delivery_date = parse_date_value(
        row_data["DeliveryDate"]
    )

    contract_info = find_contract_info(
        workbook,
        row_data["VendorID"],
        delivery_date
    )

    if contract_info is not None:
        if not row_data["VendorName"]:
            row_data["VendorName"] = contract_info["VendorName"]

        if not row_data["ItemName"]:
            row_data["ItemName"] = contract_info["ItemName"]

        if not row_data["ItemCode"]:
            row_data["ItemCode"] = contract_info["ItemCode"]

    base_rate = (
        contract_info.get("BaseRate")
        if contract_info is not None
        else None
    )

    row_data["BaseRate"] = fallback(
        row_data["BaseRate"],
        base_rate
    )

    row_data["AdjustedRate"] = fallback(
        row_data["AdjustedRate"],
        row_data["BaseRate"]
    )

    if row_data["DrcWeight"] is not None and row_data["AdjustedRate"] is not None:
        row_data["TaxableAmount"] = fallback(
            row_data["TaxableAmount"],
            round_money(row_data["DrcWeight"] * row_data["AdjustedRate"])
        )

    if row_data["TaxableAmount"] is not None and gst_percent is not None:
        row_data["GstAmount"] = fallback(
            row_data["GstAmount"],
            round_money(row_data["TaxableAmount"] * gst_percent / 100)
        )

    if row_data["TaxableAmount"] is not None:
        row_data["GrossAmount"] = fallback(
            row_data["GrossAmount"],
            round_money(
                row_data["TaxableAmount"] +
                (row_data["GstAmount"] or 0) +
                unloading_charge
            )
        )

    if row_data["TaxableAmount"] is not None and tds_percent is not None:
        row_data["TdsAmount"] = fallback(
            row_data["TdsAmount"],
            round_money(row_data["TaxableAmount"] * tds_percent / 100)
        )

    if row_data["GrossAmount"] is not None:
        row_data["NetPayable"] = fallback(
            row_data["NetPayable"],
            round_money(
                row_data["GrossAmount"] -
                (row_data["TdsAmount"] or 0)
            )
        )

    return row_data


# ============================================================
# READ PURCHASE DATA
# ============================================================

def read_purchase_data():
    database_path = get_database_path()

    workbook = load_workbook(
        database_path,
        data_only=True
    )

    if PURCHASE_SHEET_NAME not in workbook.sheetnames:
        raise ValueError(
            f"Sheet not found: {PURCHASE_SHEET_NAME}"
        )

    sheet = workbook[PURCHASE_SHEET_NAME]
    purchase_data = []

    for row in range(START_ROW, sheet.max_row + 1):
        vendor_name = cell_value(sheet, COL_VENDOR_NAME, row)
        vendor_id = cell_value(sheet, COL_VENDOR_ID, row)
        invoice_number = cell_value(sheet, COL_INVOICE_NUMBER, row)

        if vendor_name is None and vendor_id is None and invoice_number is None:
            continue

        row_data = {
            "RowNumber": row,
            "VendorName": text_value(vendor_name),
            "VendorID": text_value(vendor_id),
            "ItemName": text_value(cell_value(sheet, COL_ITEM_NAME, row)),
            "ItemCode": text_value(cell_value(sheet, COL_ITEM_CODE, row)),
            "InvoiceNumber": text_value(invoice_number),
            "PurchaseOrderDate": date_value(cell_value(sheet, COL_PURCHASE_ORDER, row)),
            "DeliveryDate": date_value(cell_value(sheet, COL_DELIVERY, row)),
            "InvoiceWeight": number_value(cell_value(sheet, COL_INVOICE_WEIGHT, row)),
            "BeforeUnloading": number_value(cell_value(sheet, COL_BEFORE_UNLOADING, row)),
            "CarrierWeight": number_value(cell_value(sheet, COL_CARRIER_WEIGHT, row)),
            "ReceivedWeight": number_value(cell_value(sheet, COL_RECEIVED_WEIGHT, row)),
            "NoOfBags": int_value(cell_value(sheet, COL_NO_OF_BAGS, row)),
            "NetWeight": number_value(cell_value(sheet, COL_NET_WEIGHT, row)),
            "CalculatedDrc": number_value(cell_value(sheet, COL_CALCULATED_DRC, row)),
            "DrcWeight": number_value(cell_value(sheet, COL_DRC_WEIGHT, row)),
            "BaseRate": number_value(cell_value(sheet, COL_BASE_RATE, row)),
            "AdjustedRate": number_value(cell_value(sheet, COL_ADJUSTED_RATE, row)),
            "GstPercent": number_value(cell_value(sheet, COL_GST_PERCENT, row)),
            "Tds194QPercent": number_value(cell_value(sheet, COL_TDS_194Q_PERCENT, row)),
            "TaxableAmount": number_value(cell_value(sheet, COL_TAXABLE_AMOUNT, row)),
            "UnloadingCharge": number_value(cell_value(sheet, COL_UNLOADING_CHARGE, row)),
            "GstAmount": number_value(cell_value(sheet, COL_GST_AMOUNT, row)),
            "GrossAmount": number_value(cell_value(sheet, COL_GROSS_AMOUNT, row)),
            "TdsAmount": number_value(cell_value(sheet, COL_TDS_AMOUNT, row)),
            "NetPayable": number_value(cell_value(sheet, COL_NET_PAYABLE, row)),
        }

        purchase_data.append(
            apply_calculated_fallbacks(
                workbook,
                row_data
            )
        )

    return purchase_data


# ============================================================
# MAIN
# ============================================================

try:
    print(
        json.dumps(
            read_purchase_data(),
            ensure_ascii=False
        )
    )

except Exception as error:
    print(f"ERROR: {error}", file=sys.stderr)
    sys.exit(1)
