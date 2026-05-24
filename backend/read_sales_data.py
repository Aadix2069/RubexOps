import json
import os
import sys
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import load_workbook


# ============================================================
# CONFIG
# ============================================================

CONFIG_PATH = (
    Path.home()
    / "Documents"
    / "RubexOps"
    / "database_config.json"
)

SALES_SHEET_NAME = "sales"
START_ROW = 5


# ============================================================
# HELPERS
# ============================================================

def fail(message):
    print(f"ERROR: {message}", file=sys.stderr)
    sys.exit(1)


def safe_string(value):
    if value is None:
        return ""

    return str(value).strip()


def safe_number(value):
    if value is None:
        return 0

    try:
        if isinstance(value, str):
            text = value.strip().replace(",", "").replace("%", "")

            if not text:
                return 0

            return float(Decimal(text))

        return float(value)
    except (InvalidOperation, ValueError, TypeError):
        return 0


def parse_date(value):
    if value is None:
        return None

    if isinstance(value, datetime):
        return value.date()

    if isinstance(value, date):
        return value

    text = str(value).strip()

    if not text:
        return None

    for date_format in ("%d-%m-%Y", "%d/%m/%Y", "%Y-%m-%d", "%m/%d/%Y"):
        try:
            return datetime.strptime(text, date_format).date()
        except ValueError:
            continue

    return None


def safe_date(value):
    parsed_date = parse_date(value)

    if parsed_date is None:
        return safe_string(value)

    return parsed_date.strftime("%d-%m-%Y")


def get_database_path():
    if not CONFIG_PATH.exists():
        fail("database_config.json not found.")

    try:
        with open(CONFIG_PATH, "r", encoding="utf-8") as file:
            config = json.load(file)
    except json.JSONDecodeError:
        fail("Invalid JSON inside database_config.json.")

    database_path = (
        config.get("database_path")
        or config.get("excel_path")
        or config.get("path")
        or ""
    ).strip()

    if not database_path:
        fail("Database path is empty.")

    if not os.path.exists(database_path):
        fail("Database file not found.")

    return database_path


# ============================================================
# READ SALES DATA
# ============================================================

def read_sales_data():
    database_path = get_database_path()

    try:
        workbook = load_workbook(database_path, data_only=True, read_only=True)
    except PermissionError:
        fail("Close Excel workbook before continuing.")

    try:
        if SALES_SHEET_NAME not in workbook.sheetnames:
            fail(f"Sheet not found: {SALES_SHEET_NAME}")

        sheet = workbook[SALES_SHEET_NAME]
        records = []

        for row in range(START_ROW, sheet.max_row + 1):
            customer_id = safe_string(sheet[f"C{row}"].value)
            invoice_number = safe_string(sheet[f"F{row}"].value)

            if not customer_id and not invoice_number:
                continue

            records.append(
                {
                    "row": row,
                    "customer_name": safe_string(sheet[f"B{row}"].value),
                    "customer_id": customer_id,
                    "item_name": safe_string(sheet[f"D{row}"].value),
                    "item_code": safe_string(sheet[f"E{row}"].value),
                    "invoice_number": invoice_number,
                    "sales_order": safe_date(sheet[f"G{row}"].value),
                    "dispatch": safe_date(sheet[f"H{row}"].value),
                    "weight": safe_number(sheet[f"I{row}"].value),
                    "base_rate": safe_number(sheet[f"J{row}"].value),
                    "gst_percent": safe_number(sheet[f"K{row}"].value),
                    "tcs_percent": safe_number(sheet[f"L{row}"].value),
                    "taxable_amount": safe_number(sheet[f"M{row}"].value),
                    "gst_amount": safe_number(sheet[f"N{row}"].value),
                    "total": safe_number(sheet[f"O{row}"].value),
                    "loading_charge": safe_number(sheet[f"P{row}"].value),
                    "tcs_amount": safe_number(sheet[f"Q{row}"].value),
                    "net_receivable": safe_number(sheet[f"R{row}"].value),
                }
            )

        return records

    finally:
        workbook.close()


try:
    print(json.dumps(read_sales_data(), indent=4))
except Exception as error:
    fail(str(error))

