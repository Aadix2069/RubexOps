import json
import os
import sys
import warnings
from datetime import date, datetime
from decimal import Decimal, InvalidOperation

from openpyxl import load_workbook

warnings.filterwarnings("ignore")


# =========================================================
# CONFIG
# =========================================================

CONFIG_PATH = os.path.join(
    os.path.expanduser("~"),
    "Documents",
    "RubexOps",
    "database_config.json"
)

PCON_SHEET_NAME = "pcon"
PURCHASE_SHEET_NAME = "purchase"
START_ROW = 5


# =========================================================
# COLUMN MAP - PCON
# =========================================================

COL_VENDOR_NAME = "B"
COL_VENDOR_ID = "C"
COL_ITEM_NAME = "D"
COL_ITEM_CODE = "E"
COL_START_DATE = "F"
COL_END_DATE = "G"
COL_DAYS_REMAINING = "H"
COL_DELIVERIES_SO_FAR = "J"
COL_RECENT_DELIVERY = "K"
COL_BASE_PRICE = "L"
COL_AGREED_QTY = "M"
COL_DELIVERED_QTY = "N"
COL_REMAINING_QTY = "O"
COL_COMPLETION_PERCENT = "P"
COL_BREACH = "Q"
COL_BREACH_RESPONSIBILITY = "R"
COL_PENALTY_PERCENT = "S"
COL_REVISED_RATE = "T"
COL_PENALTY_VALUE = "U"
COL_REMEDY_DAYS = "V"
COL_REMEDY_DEADLINE = "W"
COL_REMEDY_STATUS = "X"
COL_RENEWAL_REFERENCE = "Y"


# =========================================================
# COLUMN MAP - PURCHASE FALLBACK
# =========================================================

PUR_COL_VENDOR_ID = "C"
PUR_COL_ITEM_CODE = "E"
PUR_COL_DELIVERY_DATE = "H"
PUR_COL_NET_WEIGHT = "N"
PUR_COL_DRC_WEIGHT = "P"


# =========================================================
# SAFE ERROR HANDLER
# =========================================================

def fail(message):
    sys.stderr.write(f"ERROR: {message}\n")
    sys.exit(1)


# =========================================================
# DATABASE CONFIG
# =========================================================

def get_database_path():
    if not os.path.exists(CONFIG_PATH):
        fail("database_config.json not found.")

    try:
        with open(CONFIG_PATH, "r", encoding="utf-8") as file:
            config = json.load(file)
    except json.JSONDecodeError:
        fail("Invalid JSON inside database_config.json.")
    except Exception as error:
        fail(str(error))

    database_path = (
        config.get("database_path")
        or config.get("excel_path")
        or config.get("path")
        or ""
    ).strip()

    if database_path == "":
        fail("Database path is empty.")

    if not os.path.exists(database_path):
        fail("Database file not found.")

    return database_path


# =========================================================
# SAFE VALUE FUNCTIONS
# =========================================================

def is_blank(value):
    return value is None or str(value).strip() == ""


def safe_string(value):
    if value is None:
        return ""

    return str(value).strip()


def number_or_none(value):
    if value is None:
        return None

    if isinstance(value, bool):
        return None

    try:
        if isinstance(value, str):
            text = (
                value.strip()
                .replace(",", "")
                .replace("%", "")
            )

            if text == "":
                return None

            return float(Decimal(text))

        return float(value)

    except (ValueError, TypeError, InvalidOperation):
        return None


def safe_number(value):
    number = number_or_none(value)

    if number is None:
        return 0

    return round(number, 4)


def safe_percent(value):
    number = number_or_none(value)

    if number is None:
        return 0

    if 0 < abs(number) <= 1:
        number = number * 100

    return round(number, 4)


def parse_date(value):
    if value is None:
        return None

    if isinstance(value, datetime):
        return value.date()

    if isinstance(value, date):
        return value

    text = str(value).strip()

    if text == "":
        return None

    for date_format in (
        "%d-%m-%Y",
        "%d/%m/%Y",
        "%Y-%m-%d",
        "%m/%d/%Y",
    ):
        try:
            return datetime.strptime(
                text,
                date_format
            ).date()
        except ValueError:
            continue

    return None


def safe_date(value):
    parsed_date = parse_date(value)

    if parsed_date is None:
        return safe_string(value)

    return parsed_date.strftime("%d-%m-%Y")


def is_yes(value):
    text = safe_string(value).upper()

    return text in {
        "YES",
        "Y",
        "TRUE",
        "1",
        "BREACH",
        "BREACHED"
    }


# =========================================================
# PURCHASE FALLBACK CALCULATION
# =========================================================

def build_purchase_records(workbook):
    if PURCHASE_SHEET_NAME not in workbook.sheetnames:
        return []

    sheet = workbook[PURCHASE_SHEET_NAME]
    records = []

    for row in range(START_ROW, sheet.max_row + 1):
        vendor_id = safe_string(
            sheet[f"{PUR_COL_VENDOR_ID}{row}"].value
        )

        item_code = safe_string(
            sheet[f"{PUR_COL_ITEM_CODE}{row}"].value
        )

        delivery_date = parse_date(
            sheet[f"{PUR_COL_DELIVERY_DATE}{row}"].value
        )

        if not vendor_id or delivery_date is None:
            continue

        delivered_qty = number_or_none(
            sheet[f"{PUR_COL_DRC_WEIGHT}{row}"].value
        )

        if delivered_qty is None:
            delivered_qty = number_or_none(
                sheet[f"{PUR_COL_NET_WEIGHT}{row}"].value
            )

        if delivered_qty is None:
            delivered_qty = 0

        records.append(
            {
                "vendor_id": vendor_id.casefold(),
                "item_code": item_code.casefold(),
                "delivery_date": delivery_date,
                "delivered_qty": delivered_qty
            }
        )

    return records


def calculate_purchase_totals(
    purchase_records,
    vendor_id,
    item_code,
    start_date,
    end_date
):
    if not vendor_id or start_date is None or end_date is None:
        return {
            "deliveries_so_far": 0,
            "recent_delivery": None,
            "delivered_qty": 0
        }

    vendor_key = vendor_id.casefold()
    item_key = item_code.casefold()

    matching_records = []

    for record in purchase_records:
        if record["vendor_id"] != vendor_key:
            continue

        if item_key and record["item_code"] and record["item_code"] != item_key:
            continue

        if start_date <= record["delivery_date"] <= end_date:
            matching_records.append(record)

    delivered_qty = sum(
        record["delivered_qty"]
        for record in matching_records
    )

    recent_delivery = (
        max(record["delivery_date"] for record in matching_records)
        if matching_records
        else None
    )

    return {
        "deliveries_so_far": len(matching_records),
        "recent_delivery": recent_delivery,
        "delivered_qty": delivered_qty
    }


# =========================================================
# STATUS ENGINE
# =========================================================

def calculate_status(
    start_date,
    end_date,
    remaining_qty,
    breach_detected,
    renewal_reference
):
    today = date.today()

    if safe_string(renewal_reference) != "":
        return "Expired", "#F59E0B"

    if remaining_qty <= 0:
        if end_date is not None and today > end_date:
            return "Expired", "#F59E0B"

        return "Completed", "#8B5CF6"

    if start_date is not None and today < start_date:
        return "Upcoming", "#3B82F6"

    if end_date is not None and today > end_date and remaining_qty > 0:
        return "Violated", "#EF4444"

    if breach_detected and remaining_qty > 0:
        return "Violated", "#EF4444"

    return "Active", "#10B981"


# =========================================================
# MAIN READ
# =========================================================

database_path = get_database_path()

try:
    workbook = load_workbook(
        database_path,
        data_only=True
    )
except PermissionError:
    fail("Close Excel workbook before continuing.")
except Exception as error:
    fail(f"Failed to load workbook.\n{str(error)}")

try:
    if PCON_SHEET_NAME not in workbook.sheetnames:
        fail("Sheet not found.")

    sheet = workbook[PCON_SHEET_NAME]
    purchase_records = build_purchase_records(workbook)
    contracts = []

    for row in range(START_ROW, sheet.max_row + 1):
        vendor_name = safe_string(
            sheet[f"{COL_VENDOR_NAME}{row}"].value
        )

        vendor_id = safe_string(
            sheet[f"{COL_VENDOR_ID}{row}"].value
        )

        if vendor_name == "" and vendor_id == "":
            continue

        item_code = safe_string(
            sheet[f"{COL_ITEM_CODE}{row}"].value
        )

        start_date = parse_date(
            sheet[f"{COL_START_DATE}{row}"].value
        )

        end_date = parse_date(
            sheet[f"{COL_END_DATE}{row}"].value
        )

        agreed_qty = safe_number(
            sheet[f"{COL_AGREED_QTY}{row}"].value
        )

        excel_delivered_qty = number_or_none(
            sheet[f"{COL_DELIVERED_QTY}{row}"].value
        )

        excel_remaining_qty = number_or_none(
            sheet[f"{COL_REMAINING_QTY}{row}"].value
        )

        excel_completion_percent = number_or_none(
            sheet[f"{COL_COMPLETION_PERCENT}{row}"].value
        )

        excel_deliveries_so_far = number_or_none(
            sheet[f"{COL_DELIVERIES_SO_FAR}{row}"].value
        )

        purchase_totals = calculate_purchase_totals(
            purchase_records,
            vendor_id,
            item_code,
            start_date,
            end_date
        )

        delivered_qty = (
            excel_delivered_qty
            if excel_delivered_qty is not None
            else 0
        )

        if purchase_totals["delivered_qty"] > 0 and delivered_qty == 0:
            delivered_qty = purchase_totals["delivered_qty"]

        calculated_remaining_qty = max(
            agreed_qty - delivered_qty,
            0
        )

        if excel_remaining_qty is None:
            remaining_qty = calculated_remaining_qty
        elif (
            purchase_totals["delivered_qty"] > 0
            and abs(excel_remaining_qty - calculated_remaining_qty) > 0.01
        ):
            remaining_qty = calculated_remaining_qty
        else:
            remaining_qty = excel_remaining_qty

        calculated_completion_percent = (
            min((delivered_qty / agreed_qty) * 100, 100)
            if agreed_qty > 0
            else 0
        )

        if excel_completion_percent is None:
            completion_percent = calculated_completion_percent
        else:
            completion_percent = safe_percent(excel_completion_percent)

            if (
                purchase_totals["delivered_qty"] > 0
                and completion_percent == 0
                and calculated_completion_percent > 0
            ):
                completion_percent = calculated_completion_percent

        deliveries_so_far = (
            int(excel_deliveries_so_far)
            if excel_deliveries_so_far is not None
            else purchase_totals["deliveries_so_far"]
        )

        if deliveries_so_far == 0 and purchase_totals["deliveries_so_far"] > 0:
            deliveries_so_far = purchase_totals["deliveries_so_far"]

        recent_delivery_value = sheet[f"{COL_RECENT_DELIVERY}{row}"].value
        recent_delivery = safe_date(recent_delivery_value)

        if recent_delivery == "" and purchase_totals["recent_delivery"] is not None:
            recent_delivery = purchase_totals["recent_delivery"].strftime(
                "%d-%m-%Y"
            )

        breach_value = safe_string(
            sheet[f"{COL_BREACH}{row}"].value
        )

        breach_detected = is_yes(breach_value)

        renewal_reference = safe_string(
            sheet[f"{COL_RENEWAL_REFERENCE}{row}"].value
        )

        dynamic_status, status_color = calculate_status(
            start_date,
            end_date,
            remaining_qty,
            breach_detected,
            renewal_reference
        )

        is_expired = dynamic_status == "Expired"

        contract = {
            "row": row,

            "vendor_name": vendor_name,
            "vendor_id": vendor_id,
            "item_name": safe_string(sheet[f"{COL_ITEM_NAME}{row}"].value),
            "item_code": item_code,

            "start_date": safe_date(sheet[f"{COL_START_DATE}{row}"].value),
            "end_date": safe_date(sheet[f"{COL_END_DATE}{row}"].value),
            "days_remaining": safe_string(sheet[f"{COL_DAYS_REMAINING}{row}"].value),

            "deliveries_so_far": str(deliveries_so_far),
            "recent_delivery": recent_delivery,

            "base_price": safe_number(sheet[f"{COL_BASE_PRICE}{row}"].value),
            "agreed_qty": agreed_qty,
            "delivered_qty": round(delivered_qty, 4),
            "remaining_qty": round(remaining_qty, 4),
            "completion_percent": round(completion_percent, 4),

            "breach": "YES" if breach_detected else "NO",
            "breach_responsibility": safe_string(
                sheet[f"{COL_BREACH_RESPONSIBILITY}{row}"].value
            ),

            "penalty_percent": safe_percent(
                sheet[f"{COL_PENALTY_PERCENT}{row}"].value
            ),
            "revised_rate": safe_number(sheet[f"{COL_REVISED_RATE}{row}"].value),
            "penalty_value": safe_number(sheet[f"{COL_PENALTY_VALUE}{row}"].value),

            "remedy_days": safe_number(sheet[f"{COL_REMEDY_DAYS}{row}"].value),
            "remedy_deadline": safe_date(sheet[f"{COL_REMEDY_DEADLINE}{row}"].value),
            "remedy_status": safe_string(sheet[f"{COL_REMEDY_STATUS}{row}"].value),

            "renewal_reference": renewal_reference,

            "status": dynamic_status,
            "dynamic_status": dynamic_status,
            "status_color": status_color,

            "is_expired": is_expired,
            "can_edit": not is_expired,
            "can_renew": (
                not is_expired
                and dynamic_status in ["Completed", "Violated"]
                and renewal_reference == ""
            ),
            "is_breach_detected": breach_detected,
            "show_responsibility": breach_detected,
            "show_breach_panel": breach_detected
        }

        contracts.append(contract)

    contracts.sort(
        key=lambda item: (
            item["is_expired"],
            {
                "Active": 0,
                "Upcoming": 1,
                "Violated": 2,
                "Completed": 3,
                "Expired": 4
            }.get(item["dynamic_status"], 5),
            item["vendor_name"].lower()
        )
    )

    print(
        json.dumps(
            contracts,
            indent=4
        )
    )

except Exception as error:
    fail(f"Failed to read contracts.\n{str(error)}")

finally:
    workbook.close()