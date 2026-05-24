import json
import os
import sys
from datetime import datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import load_workbook


# ============================================================
# CONFIG
# ============================================================

CONFIG_PATH = (
    Path.home()
    / "Documents"
    / "RubexOps"
    / "database_config.json"
)

PURCHASE_SHEET_NAME = "purchase"
PRODUCTION_SHEET_NAME = "production"
START_ROW = 5


# ============================================================
# ERROR HANDLER
# ============================================================

def fail(message):
    print(f"ERROR: {message}", file=sys.stderr)
    sys.exit(1)


# ============================================================
# HELPERS
# ============================================================

def safe_string(value):
    if value is None:
        return ""

    return str(value).strip()


def parse_decimal(value, field_name):
    text = safe_string(value).replace(",", "").replace("%", "")

    if not text:
        fail(f"{field_name} is required.")

    try:
        return float(Decimal(text))
    except (InvalidOperation, ValueError):
        fail(f"{field_name} must be numeric.")


def number_or_zero(value):
    if value is None:
        return 0.0

    try:
        if isinstance(value, str):
            text = value.strip().replace(",", "").replace("%", "")

            if not text:
                return 0.0

            return float(Decimal(text))

        return float(value)
    except (InvalidOperation, ValueError, TypeError):
        return 0.0


def parse_date(value, field_name):
    text = safe_string(value)

    if not text:
        fail(f"{field_name} is required.")

    try:
        return datetime.strptime(text, "%d-%m-%Y")
    except ValueError:
        fail(f"{field_name} must be in dd-MM-yyyy format.")


def get_database_path():
    if not CONFIG_PATH.exists():
        fail("database_config.json not found.")

    try:
        with open(CONFIG_PATH, "r", encoding="utf-8") as file:
            config = json.load(file)
    except json.JSONDecodeError:
        fail("Invalid JSON inside database_config.json.")

    database_path = (
        config.get("database_path")
        or config.get("excel_path")
        or config.get("path")
        or ""
    ).strip()

    if not database_path:
        fail("Database path is empty.")

    if not os.path.exists(database_path):
        fail("Database file not found.")

    return database_path


def stock_key(supplier_id, invoice_number, raw_material_code):
    return (
        safe_string(supplier_id).casefold(),
        safe_string(invoice_number).casefold(),
        safe_string(raw_material_code).casefold(),
    )


def get_purchase_quantity(sheet, row):
    net_weight = number_or_zero(sheet[f"N{row}"].value)

    if net_weight > 0:
        return net_weight

    drc_weight = number_or_zero(sheet[f"P{row}"].value)

    if drc_weight > 0:
        return drc_weight

    return number_or_zero(sheet[f"I{row}"].value)


def calculate_available_quantity(workbook, supplier_id, invoice_number, raw_material_code, exclude_row):
    if PURCHASE_SHEET_NAME not in workbook.sheetnames:
        fail(f"Sheet not found: {PURCHASE_SHEET_NAME}")

    key = stock_key(supplier_id, invoice_number, raw_material_code)
    purchased = 0.0
    consumed = 0.0
    purchase_sheet = workbook[PURCHASE_SHEET_NAME]
    production_sheet = workbook[PRODUCTION_SHEET_NAME]

    for row in range(START_ROW, purchase_sheet.max_row + 1):
        row_key = stock_key(
            purchase_sheet[f"C{row}"].value,
            purchase_sheet[f"F{row}"].value,
            purchase_sheet[f"E{row}"].value,
        )

        if row_key == key:
            purchased += get_purchase_quantity(purchase_sheet, row)

    for row in range(START_ROW, production_sheet.max_row + 1):
        if row == exclude_row:
            continue

        row_key = stock_key(
            production_sheet[f"E{row}"].value,
            production_sheet[f"F{row}"].value,
            production_sheet[f"H{row}"].value,
        )

        if row_key == key:
            consumed += number_or_zero(production_sheet[f"L{row}"].value)

    return round(max(purchased - consumed, 0.0), 4)


# ============================================================
# UPDATE RECORD
# ============================================================

def update_record(args):
    if len(args) != 13:
        fail(
            "Expected 13 arguments: row, production date, supplier name, supplier id, "
            "invoice number, raw material, raw material code, finished product, "
            "finished product code, input weight, output weight, initial DRC, actual DRC."
        )

    try:
        row_number = int(args[0])
    except ValueError:
        fail("Invalid production row.")

    production_date = parse_date(args[1], "Production Date")
    supplier_name = safe_string(args[2])
    supplier_id = safe_string(args[3])
    invoice_number = safe_string(args[4])
    raw_material = safe_string(args[5])
    raw_material_code = safe_string(args[6])
    finished_product = safe_string(args[7])
    finished_product_code = safe_string(args[8])
    input_weight = parse_decimal(args[9], "Input Weight")
    output_weight = parse_decimal(args[10], "Output Weight")
    initial_drc = parse_decimal(args[11], "Initial DRC")
    actual_drc = parse_decimal(args[12], "Actual DRC")

    if input_weight <= 0:
        fail("Input Weight must be greater than zero.")

    if output_weight < 0:
        fail("Output Weight cannot be negative.")

    if output_weight > input_weight:
        fail("Output Weight cannot exceed Input Weight.")

    if initial_drc < 0 or initial_drc > 100:
        fail("Initial DRC must be between 0 and 100.")

    if actual_drc < 0 or actual_drc > 100:
        fail("Actual DRC must be between 0 and 100.")

    database_path = get_database_path()

    try:
        workbook = load_workbook(database_path, data_only=False)
    except PermissionError:
        fail("Close Excel workbook before continuing.")
    except Exception as error:
        fail(f"Failed to open workbook.\n{error}")

    try:
        if PRODUCTION_SHEET_NAME not in workbook.sheetnames:
            fail(f"Sheet not found: {PRODUCTION_SHEET_NAME}")

        sheet = workbook[PRODUCTION_SHEET_NAME]

        if row_number < START_ROW or row_number > sheet.max_row:
            fail("Production row does not exist.")

        if not safe_string(sheet[f"B{row_number}"].value):
            fail("No production record exists in selected row.")

        quantity_available = calculate_available_quantity(
            workbook,
            supplier_id,
            invoice_number,
            raw_material_code,
            exclude_row=row_number,
        )

        if input_weight > quantity_available:
            fail(
                f"Input Weight exceeds available quantity. "
                f"Available quantity is {quantity_available:0.##}."
            )

        sheet[f"C{row_number}"] = production_date
        sheet[f"C{row_number}"].number_format = "dd-mm-yyyy"
        sheet[f"D{row_number}"] = supplier_name
        sheet[f"E{row_number}"] = supplier_id
        sheet[f"F{row_number}"] = invoice_number
        sheet[f"G{row_number}"] = raw_material
        sheet[f"H{row_number}"] = raw_material_code
        sheet[f"I{row_number}"] = finished_product
        sheet[f"J{row_number}"] = finished_product_code
        sheet[f"K{row_number}"] = quantity_available
        sheet[f"L{row_number}"] = input_weight
        sheet[f"M{row_number}"] = output_weight
        sheet[f"N{row_number}"] = f"=K{row_number}-L{row_number}"
        sheet[f"O{row_number}"] = f"=L{row_number}-M{row_number}"
        sheet[f"P{row_number}"] = initial_drc
        sheet[f"Q{row_number}"] = actual_drc
        sheet[f"R{row_number}"] = f"=Q{row_number}-P{row_number}"

        workbook.calculation.fullCalcOnLoad = True
        workbook.calculation.forceFullCalc = True
        workbook.save(database_path)

        return "Production record updated successfully."

    except PermissionError:
        fail("Cannot save workbook. Close Excel file first.")
    finally:
        workbook.close()


# ============================================================
# MAIN
# ============================================================

try:
    print(update_record(sys.argv[1:]))
except Exception as error:
    fail(str(error))

